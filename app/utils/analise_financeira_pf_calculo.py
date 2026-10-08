# app/utils/analise_financeira_pf_calculo.py
"""
Motor de cálculo do Fluxo de Caixa Simulado - Análise Financeira PF.

Reproduz a aba "Fluxo 1" da planilha da SUFIN, célula a célula:

  J  = PZ_EXECUCAO_MESES                       (mês da consolidação; Excel J7)
  L  = J + PZ_PERMANENCIA_ESTOQUE_MESES         (mês da venda;        Excel L7)
  VA = VR_LAUDO_AVALIACAO do contrato           (Excel O13)

  Para cada NU_MES n = 0 .. L  (Excel colunas A..I):
    ANO_MES             = data de referência + n meses (1º dia do mês)
    VR_DESP_MANUT       (col. C)
        n = J + 12      -> -VA * PC_DESP_MANUT_INU_PRI_ANO      (L4)
        n = J + 24      ->  0  (Memoria C42 = 0%, sem coluna na tabela) (L5)
        n = J + 36      ->  0  (Memoria C42 = 0%, sem coluna na tabela) (L6)
        n = L           -> -VA * PC_DESP_MANUT_INU_6_MESES       (L7)
        (a ordem dos testes é a mesma do SE() aninhado do Excel)
    VR_CUSTO_MANUT      (col. D)
        n <= J          -> -VR_CUSTO_CIPF
        n >  J          -> -(VR_TARIFA_ADM_IMOVEIS + VR_CUSTO_INU)
    VR_DESP_CONSOL_PROP (col. E)  n = J -> -VR_DESP_MEDIA_EXEC_JUD
    VR_DEB_PROPTERREM   (col. F)  n = J -> -VR_DEBITOS_PROPTERREM do contrato
    VR_VENDA            (col. G)  n = L -> VA * (1 - PC_DESP_DESCONTO_VENDA)
    VR_TOTAL            (col. H)  soma de C..G
    VR_PRESENTE         (col. I)  VR_TOTAL / (1 + taxa) ^ (n / 12)
        taxa = CUSTO_MEDIO / 100 da view FIN_VW036 no ANO_MES da linha;
        mês sem taxa na view repete a taxa do mês anterior (IFERROR do Excel).

  VPL do contrato = soma de VR_PRESENTE (Excel I7).

Simulação de Equilíbrio (abas "Meta" da planilha) - VPL negativo:
  O VPL é linear nos débitos propter rem (só entram no mês J), então o valor
  que zera o VPL é obtido direto, sem "Atingir Meta":
      P* = TOTAL_J sem propter + VPL dos outros meses * (1 + taxa_J)^(J/12)
  - P* > 0  -> NEGATIVO REVERSÍVEL (reduzir os débitos até P* zera o VPL)
  - P* <= 0 -> NEGATIVO IRREVERSÍVEL (mesmo com débitos = 0 o VPL é negativo)
  Gravada em FIN_TB037. Cada contrato tem um único cálculo vigente:
  recalcular apaga o fluxo e a simulação anteriores do contrato.

Valores nominais em Decimal; o desconto (expoente fracionário) em float,
como o Excel. Cada valor gravado é arredondado em 2 casas (ROUND_HALF_UP).

Compatível com Python 3.9 e 3.12.
"""
from app import db
from sqlalchemy import text
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP, ROUND_DOWN
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


TB_PARAMETROS = '[BDG].[FIN_TB034_ANALISE_FINANCEIRA_PF_PARAMETROS]'
TB_CONTRATOS = '[BDG].[FIN_TB035_ANALISE_FINANCEIRA_PF_CONTRATOS]'
TB_FLUXOS = '[BDG].[FIN_TB036_ANALISE_FINANCEIRA_PF_FLUXOS]'
TB_SIMULACAO = '[BDG].[FIN_TB037_ANALISE_FINANCEIRA_PF_SIMULACAO]'
VW_CUSTO_MEDIO = '[BDG].[FIN_VW036_ANALISE_FINANCEIRA_PF_CUSTO_MEDIO]'

CENTAVO = Decimal('0.01')

MESES_ABREV = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun',
               'jul', 'ago', 'set', 'out', 'nov', 'dez']

# Colunas de valor do fluxo, na ordem da planilha (C..I)
COLUNAS_FLUXO = [
    ('VR_DESP_MANUT', 'Despesas de Manutenção'),
    ('VR_CUSTO_MANUT', 'Custo Manutenção'),
    ('VR_DESP_CONSOL_PROP', 'Despesas de Consolidação de Propriedade'),
    ('VR_DEB_PROPTERREM', 'Débitos Propter Rem'),
    ('VR_VENDA', 'Venda Imóvel Executado'),
    ('VR_TOTAL', 'TOTAL'),
    ('VR_PRESENTE', 'Valor Presente'),
]


class ErroCalculo(Exception):
    """Erro de regra de negócio no cálculo (mensagem vai para a tela)."""
    pass


# =========================================================================
# UTILITÁRIOS
# =========================================================================
def _dec(valor):
    if valor is None:
        return Decimal('0')
    return Decimal(str(valor))


def _r2(valor):
    return _dec(valor).quantize(CENTAVO, rounding=ROUND_HALF_UP)


def _para_date(valor):
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.date()
    return valor


def somar_meses(data_base, meses):
    """Equivalente ao EDATE do Excel, sempre no 1º dia do mês."""
    total = data_base.year * 12 + (data_base.month - 1) + meses
    return date(total // 12, total % 12 + 1, 1)


def chave_ano_mes(data_ref):
    """date(2026, 9, 1) -> 202609"""
    return data_ref.year * 100 + data_ref.month


def mes_ano_texto(data_ref):
    """date(2026, 9, 1) -> 'set/2026' (mesmo formato da planilha)."""
    if not data_ref:
        return ''
    return f'{MESES_ABREV[data_ref.month - 1]}/{data_ref.year}'


def fmt_br(valor, casas=2):
    if valor is None:
        return ''
    fmt = '{:,.' + str(casas) + 'f}'
    return fmt.format(float(valor)).replace(',', 'X').replace('.', ',').replace('X', '.')


def formatar_contrato(nu_contrato):
    """Contrato exibido só com os dígitos, sem traços: 101600101095."""
    if nu_contrato is None:
        return ''
    txt = str(nu_contrato).strip()
    if txt.endswith('.0'):
        txt = txt[:-2]
    return txt


# Memo: a tabela guarda só "número/ano" (ex.: 239/2026); o restante é
# completado pelo aplicativo na exibição e no Excel.
MEMO_PREFIXO = 'Memorando SEI nº '
MEMO_SUFIXO = '/Gecoc/Sucre/Diope'


# Memorando de saída (Gefin): também só "número/ano" na tabela (MEMO_GEFIN)
MEMO_GEFIN_SUFIXO = '/Gefin/Sufin/Difin'


def formatar_memo_gefin(memo):
    """'762/2026' -> 'Memorando SEI nº 762/2026/Gefin/Sufin/Difin' ('' se vazio)."""
    txt = str(memo or '').strip()
    if not txt:
        return ''
    return f'{MEMO_PREFIXO}{txt}{MEMO_GEFIN_SUFIXO}'


def formatar_memo(memo):
    """'239/2026' -> 'Memorando SEI nº 239/2026/Gecoc/Sucre/Diope' ('' se vazio)."""
    txt = str(memo or '').strip()
    if not txt:
        return ''
    return f'{MEMO_PREFIXO}{txt}{MEMO_SUFIXO}'


def contrato_texto(nu_contrato):
    """Decimal('101600101095') -> '101600101095'"""
    txt = str(nu_contrato).strip()
    return txt[:-2] if txt.endswith('.0') else txt


# =========================================================================
# LEITURA DAS FONTES
# =========================================================================
def obter_parametros_na_data(data_ref):
    """
    Linha de FIN_TB034 em vigor na data informada:
      DT_INI_VIGENCIA <= data  E  (DT_FIM_VIGENCIA IS NULL OU DT_FIM_VIGENCIA >= data)
    Assim, no dia de uma edição ainda vale a versão antiga (que termina hoje);
    a nova só entra no dia seguinte, como definido na regra de vigência.
    """
    row = db.session.execute(text(f"""
        SELECT TOP 1
               [DT_INI_VIGENCIA], [DT_FIM_VIGENCIA],
               [VR_CUSTO_CIPF], [VR_CUSTO_INU], [VR_TARIFA_ADM_IMOVEIS],
               [PC_DESP_MANUT_INU_PRI_ANO], [PC_DESP_MANUT_INU_6_MESES],
               [PC_DESP_DESCONTO_VENDA], [VR_DESP_MEDIA_EXEC_JUD],
               [PZ_EXECUCAO_ANOS], [PZ_PERMANENCIA_ESTOQUE_MESES],
               [PZ_EXECUCAO_MESES], [DT_CUSTO_DE_OPORTUNIDADE],
               [NORMATIVO_SUFIN], [NR_ATA_DIREX], [DT_ATA_DIREX]
        FROM {TB_PARAMETROS}
        WHERE [DT_INI_VIGENCIA] <= :data_ref
          AND ([DT_FIM_VIGENCIA] IS NULL OR [DT_FIM_VIGENCIA] >= :data_ref)
        ORDER BY [DT_INI_VIGENCIA] DESC
    """), {'data_ref': data_ref}).mappings().first()
    return dict(row) if row else None


def obter_taxas_custo_medio():
    """
    Lê a view FIN_VW036 e devolve {AAAAMM (int): CUSTO_MEDIO (Decimal, % a.a.)}.
    Aceita ANO_MES vindo como número, texto ou data.
    """
    rows = db.session.execute(text(f"""
        SELECT [ANO_MES], [CUSTO_MEDIO]
        FROM {VW_CUSTO_MEDIO}
    """)).fetchall()

    taxas = {}
    for ano_mes, custo in rows:
        if ano_mes is None or custo is None:
            continue
        if isinstance(ano_mes, (date, datetime)):
            chave = ano_mes.year * 100 + ano_mes.month
        else:
            txt = str(ano_mes).strip().replace('-', '').replace('/', '')
            if txt.endswith('.0'):
                txt = txt[:-2]
            if not txt.isdigit() or len(txt) < 6:
                continue
            chave = int(txt[:6])
        taxas[chave] = _dec(custo)
    return taxas


def listar_contratos_cadastrados():
    rows = db.session.execute(text(f"""
        SELECT [NU_CONTRATO], [NO_MUTUARIO],
               [VR_DEBITOS_PROPTERREM], [VR_LAUDO_AVALIACAO]
        FROM {TB_CONTRATOS}
        ORDER BY [NO_MUTUARIO], [NU_CONTRATO]
    """)).mappings().all()
    return [dict(r) for r in rows]


def obter_contratos(lista_nu):
    """Busca os contratos escolhidos na FIN_TB035 (lista de int)."""
    if not lista_nu:
        return []
    params = {}
    marcadores = []
    for i, nu in enumerate(lista_nu):
        params[f'nu{i}'] = nu
        marcadores.append(f':nu{i}')
    rows = db.session.execute(text(f"""
        SELECT [NU_CONTRATO], [NO_MUTUARIO],
               [VR_DEBITOS_PROPTERREM], [VR_LAUDO_AVALIACAO]
        FROM {TB_CONTRATOS}
        WHERE [NU_CONTRATO] IN ({', '.join(marcadores)})
        ORDER BY [NO_MUTUARIO], [NU_CONTRATO]
    """), params).mappings().all()
    return [dict(r) for r in rows]


# =========================================================================
# CÁLCULO (função pura: não acessa banco)
# =========================================================================
def prazos_do_fluxo(parametros):
    """Retorna (J, L): mês da consolidação e mês da venda."""
    meses_exec = parametros.get('PZ_EXECUCAO_MESES')
    if meses_exec is None:
        meses_exec = int(parametros['PZ_EXECUCAO_ANOS']) * 12
    j = int(meses_exec)
    l = j + int(parametros['PZ_PERMANENCIA_ESTOQUE_MESES'])
    return j, l


def calcular_fluxo(contrato, parametros, data_referencia, taxas):
    """
    Monta o fluxo mês a mês de UM contrato.

    contrato:        dict com VR_DEBITOS_PROPTERREM e VR_LAUDO_AVALIACAO
    parametros:      dict da FIN_TB034
    data_referencia: date (1º dia do mês) = mês 0
    taxas:           {AAAAMM: CUSTO_MEDIO em % a.a.}

    Retorna (linhas, vpl). Cada linha já vem arredondada em 2 casas e com a
    taxa usada (TAXA_AA, em % a.a.) para exibição/exportação.
    """
    j, l = prazos_do_fluxo(parametros)

    va = _dec(contrato.get('VR_LAUDO_AVALIACAO'))
    debitos = _dec(contrato.get('VR_DEBITOS_PROPTERREM'))

    custo_cipf = _dec(parametros['VR_CUSTO_CIPF'])
    custo_inu = _dec(parametros['VR_CUSTO_INU'])
    tarifa_adm = _dec(parametros['VR_TARIFA_ADM_IMOVEIS'])
    pc_pri_ano = _dec(parametros['PC_DESP_MANUT_INU_PRI_ANO'])
    pc_6_meses = _dec(parametros['PC_DESP_MANUT_INU_6_MESES'])
    pc_desconto = _dec(parametros['PC_DESP_DESCONTO_VENDA'])
    desp_exec = _dec(parametros['VR_DESP_MEDIA_EXEC_JUD'])

    # Marcos anuais de manutenção (Excel L4, L5, L6)
    marco_1 = j + 12
    marco_2 = j + 24
    marco_3 = j + 36

    linhas = []
    vpl = Decimal('0')
    taxa_anterior = None

    for n in range(0, l + 1):
        mes = somar_meses(data_referencia, n)
        taxa = taxas.get(chave_ano_mes(mes))
        if taxa is None:
            if taxa_anterior is None:
                raise ErroCalculo(
                    f'Não há custo médio na FIN_VW036 para {mes_ano_texto(mes)} '
                    f'(mês 0 da data de referência).'
                )
            taxa = taxa_anterior
        taxa_anterior = taxa

        # Col. C - Despesas de manutenção (% do VA)
        if n == marco_1:
            desp_manut = -(va * pc_pri_ano)
        elif n == marco_2 or n == marco_3:
            desp_manut = Decimal('0')
        elif n == l:
            desp_manut = -(va * pc_6_meses)
        else:
            desp_manut = Decimal('0')

        # Col. D - Custo de manutenção
        if n <= j:
            custo_manut = -custo_cipf
        else:
            custo_manut = -(tarifa_adm + custo_inu)

        # Col. E e F - Consolidação e propter rem no mês J
        desp_consol = -desp_exec if n == j else Decimal('0')
        deb_propter = -debitos if n == j else Decimal('0')

        # Col. G - Venda no mês L
        venda = va * (Decimal('1') - pc_desconto) if n == l else Decimal('0')

        # Col. H - Total
        total = desp_manut + custo_manut + desp_consol + deb_propter + venda

        # Col. I - Valor presente
        fator = (1.0 + float(taxa) / 100.0) ** (n / 12.0)
        presente = Decimal(repr(float(total) / fator))

        linha = {
            'NU_MES': n,
            'ANO_MES': mes,
            'TAXA_AA': taxa,
            'VR_DESP_MANUT': _r2(desp_manut),
            'VR_CUSTO_MANUT': _r2(custo_manut),
            'VR_DESP_CONSOL_PROP': _r2(desp_consol),
            'VR_DEB_PROPTERREM': _r2(deb_propter),
            'VR_VENDA': _r2(venda),
            'VR_TOTAL': _r2(total),
            'VR_PRESENTE': _r2(presente),
            'FATOR': fator,   # (1 + taxa)^(n/12) - não é gravado, usado na simulação
        }
        vpl += linha['VR_PRESENTE']
        linhas.append(linha)

    return linhas, vpl


# =========================================================================
# SIMULAÇÃO DE EQUILÍBRIO FINANCEIRO (VPL NEGATIVO)
# =========================================================================
# Situações gravadas em IC_SITUACAO
SITUACAO_POSITIVO = 1
SITUACAO_NEGATIVO_REVERSIVEL = 2     # reduzindo os débitos propter rem o VPL zera
SITUACAO_NEGATIVO_IRREVERSIVEL = 3   # mesmo com débitos propter rem = 0 o VPL fica negativo

SITUACOES = {
    SITUACAO_POSITIVO: {
        'rotulo': 'POSITIVO', 'descricao': 'VPL positivo: a execução é vantajosa.',
        'cor': '#047857', 'fundo': '#d1fae5',
    },
    SITUACAO_NEGATIVO_REVERSIVEL: {
        'rotulo': 'NEGATIVO - REVERSÍVEL',
        'descricao': 'Reduzindo os débitos propter rem até o valor simulado, o VPL zera.',
        'cor': '#b45309', 'fundo': '#fef3c7',
    },
    SITUACAO_NEGATIVO_IRREVERSIVEL: {
        'rotulo': 'NEGATIVO - IRREVERSÍVEL',
        'descricao': 'Mesmo com os débitos propter rem zerados, o VPL continua negativo.',
        'cor': '#b91c1c', 'fundo': '#fee2e2',
    },
}


def simular_equilibrio(linhas, debitos_propterrem, mes_consolidacao):
    """
    Reproduz as abas "Meta" da planilha (Atingir Meta sobre os Débitos
    Propter Rem, célula O20), sem precisar de tentativa e erro:

    O VPL é linear nos débitos propter rem (P), que só entram no mês J:
        VPL(P) = VPL_outros + (TOTAL_J_sem_P - P) / FATOR_J
    Logo o P que zera o VPL é:
        P* = TOTAL_J_sem_P + VPL_outros * FATOR_J

    VPL_outros usa os valores presentes JÁ ARREDONDADOS (os mesmos gravados
    e exibidos), para que o VPL simulado mostrado na tela fique em ~0,00.

    Situações:
      VPL >= 0                    -> 1 POSITIVO (sem simulação)
      VPL < 0 e VPL(P=0) > 0      -> 2 REVERSÍVEL: P* entre 0 e o débito atual
      VPL < 0 e VPL(P=0) <= 0     -> 3 IRREVERSÍVEL: P simulado = 0
    """
    j = mes_consolidacao
    debitos = _r2(debitos_propterrem)
    vpl = sum((ln['VR_PRESENTE'] for ln in linhas), Decimal('0'))

    linha_j = linhas[j]
    fator_j = linha_j['FATOR']
    total_j_sem_p = linha_j['VR_TOTAL'] + debitos          # total do mês J sem o propter rem
    vpl_outros = vpl - linha_j['VR_PRESENTE']                # VP de todos os outros meses
    vp_j_sem_p = _r2(Decimal(repr(float(total_j_sem_p) / fator_j)))
    vpl_sem_p = vpl_outros + vp_j_sem_p

    resultado = {
        'IC_SITUACAO': SITUACAO_POSITIVO,
        'VR_VPL': vpl,
        'VR_DEB_PROPTERREM': debitos,
        'VR_VPL_SEM_PROPTERREM': vpl_sem_p,
        'VR_DEB_PROPTERREM_SIMULADO': None,
        'VR_REDUCAO_PROPTERREM': None,
        'VR_VPL_SIMULADO': None,
        'NU_MES_CONSOLIDACAO': j,
        'VR_TOTAL_SIMULADO': None,
        'VR_PRESENTE_SIMULADO': None,
    }

    if vpl >= 0:
        return resultado

    if vpl_sem_p <= 0:
        # Mesmo sem débitos o VPL não fica positivo
        p_sim = Decimal('0.00')
        resultado['IC_SITUACAO'] = SITUACAO_NEGATIVO_IRREVERSIVEL
    else:
        # Arredonda para baixo: o VPL simulado fica >= 0 (nunca negativo por centavos)
        p_float = float(total_j_sem_p) + float(vpl_outros) * fator_j
        p_sim = Decimal(repr(p_float)).quantize(CENTAVO, rounding=ROUND_DOWN)
        if p_sim < 0:
            p_sim = Decimal('0.00')
        resultado['IC_SITUACAO'] = SITUACAO_NEGATIVO_REVERSIVEL

    total_sim = total_j_sem_p - p_sim
    vp_sim = _r2(Decimal(repr(float(total_sim) / fator_j)))

    resultado.update({
        'VR_DEB_PROPTERREM_SIMULADO': p_sim,
        'VR_REDUCAO_PROPTERREM': debitos - p_sim,
        'VR_VPL_SIMULADO': vpl_outros + vp_sim,
        'VR_TOTAL_SIMULADO': total_sim,
        'VR_PRESENTE_SIMULADO': vp_sim,
    })
    return resultado


def montar_fluxo_meta(linhas, simulacao):
    """
    Fluxo da "Meta": igual ao fluxo do contrato, trocando no mês J os
    débitos propter rem pelo valor simulado (e o total / VP daquele mês).
    """
    if not simulacao or simulacao.get('VR_DEB_PROPTERREM_SIMULADO') is None:
        return None
    j = int(simulacao['NU_MES_CONSOLIDACAO'])
    meta = []
    for ln in linhas:
        nova = dict(ln)
        if int(ln['NU_MES']) == j:
            nova['VR_DEB_PROPTERREM'] = -_dec(simulacao['VR_DEB_PROPTERREM_SIMULADO'])
            nova['VR_TOTAL'] = _dec(simulacao['VR_TOTAL_SIMULADO'])
            nova['VR_PRESENTE'] = _dec(simulacao['VR_PRESENTE_SIMULADO'])
        meta.append(nova)
    return meta


# =========================================================================
# GRAVAÇÃO
# =========================================================================
def gravar_fluxo(dt_calculo, nu_contrato, linhas, simulacao):
    """
    Um contrato tem um único cálculo valendo: apaga TODO fluxo e simulação
    anteriores do contrato (qualquer DT_CALCULO) e grava o novo.
    """
    params_nu = {'nu': nu_contrato}
    db.session.execute(text(f"DELETE FROM {TB_FLUXOS} WHERE [NU_CONTRATO] = :nu"), params_nu)
    db.session.execute(text(f"DELETE FROM {TB_SIMULACAO} WHERE [NU_CONTRATO] = :nu"), params_nu)

    registros = [{
        'dt_calculo': dt_calculo,
        'nu': nu_contrato,
        'nu_mes': ln['NU_MES'],
        'ano_mes': ln['ANO_MES'],
        'desp_manut': ln['VR_DESP_MANUT'],
        'custo_manut': ln['VR_CUSTO_MANUT'],
        'desp_consol': ln['VR_DESP_CONSOL_PROP'],
        'deb_propter': ln['VR_DEB_PROPTERREM'],
        'venda': ln['VR_VENDA'],
        'total': ln['VR_TOTAL'],
        'presente': ln['VR_PRESENTE'],
    } for ln in linhas]

    db.session.execute(text(f"""
        INSERT INTO {TB_FLUXOS}
            ([DT_CALCULO], [NU_CONTRATO], [NU_MES], [ANO_MES],
             [VR_DESP_MANUT], [VR_CUSTO_MANUT], [VR_DESP_CONSOL_PROP],
             [VR_DEB_PROPTERREM], [VR_VENDA], [VR_TOTAL], [VR_PRESENTE])
        VALUES
            (:dt_calculo, :nu, :nu_mes, :ano_mes,
             :desp_manut, :custo_manut, :desp_consol,
             :deb_propter, :venda, :total, :presente)
    """), registros)

    db.session.execute(text(f"""
        INSERT INTO {TB_SIMULACAO}
            ([DT_CALCULO], [NU_CONTRATO], [IC_SITUACAO], [VR_VPL],
             [VR_DEB_PROPTERREM], [VR_VPL_SEM_PROPTERREM],
             [VR_DEB_PROPTERREM_SIMULADO], [VR_REDUCAO_PROPTERREM], [VR_VPL_SIMULADO],
             [NU_MES_CONSOLIDACAO], [VR_TOTAL_SIMULADO], [VR_PRESENTE_SIMULADO])
        VALUES
            (:dt_calculo, :nu, :situacao, :vpl,
             :debitos, :vpl_sem,
             :deb_sim, :reducao, :vpl_sim,
             :mes_j, :total_sim, :vp_sim)
    """), {
        'dt_calculo': dt_calculo,
        'nu': nu_contrato,
        'situacao': simulacao['IC_SITUACAO'],
        'vpl': simulacao['VR_VPL'],
        'debitos': simulacao['VR_DEB_PROPTERREM'],
        'vpl_sem': simulacao['VR_VPL_SEM_PROPTERREM'],
        'deb_sim': simulacao['VR_DEB_PROPTERREM_SIMULADO'],
        'reducao': simulacao['VR_REDUCAO_PROPTERREM'],
        'vpl_sim': simulacao['VR_VPL_SIMULADO'],
        'mes_j': simulacao['NU_MES_CONSOLIDACAO'],
        'total_sim': simulacao['VR_TOTAL_SIMULADO'],
        'vp_sim': simulacao['VR_PRESENTE_SIMULADO'],
    })


# =========================================================================
# CONSULTA DOS RESULTADOS GRAVADOS
# =========================================================================
def listar_datas_calculo():
    rows = db.session.execute(text(f"""
        SELECT [DT_CALCULO], COUNT(DISTINCT [NU_CONTRATO]) AS QTD
        FROM {TB_FLUXOS}
        GROUP BY [DT_CALCULO]
        ORDER BY [DT_CALCULO] DESC
    """)).fetchall()
    return [{'dt_calculo': _para_date(r[0]), 'qtd': int(r[1])} for r in rows]


def listar_resumo_calculo(dt_calculo=None, nu_contrato=None, filtro_contrato=None, filtro_nome=None):
    """
    Um registro por (DT_CALCULO, contrato): data de referência (ANO_MES do
    mês 0), meses do fluxo, VPL (soma de VR_PRESENTE), dados da FIN_TB035 e
    a simulação de equilíbrio (FIN_TB037).

    Filtros (todos opcionais):
      dt_calculo      -> só aquela data de cálculo (None = todas)
      nu_contrato     -> contrato exato (int)
      filtro_contrato -> parte do número do contrato (texto só com dígitos)
      filtro_nome     -> parte do nome do mutuário
    """
    condicoes = []
    params = {}
    if dt_calculo is not None:
        condicoes.append('f.[DT_CALCULO] = :dt_calculo')
        params['dt_calculo'] = dt_calculo
    if nu_contrato is not None:
        condicoes.append('f.[NU_CONTRATO] = :nu')
        params['nu'] = nu_contrato
    if filtro_contrato:
        condicoes.append("CAST(f.[NU_CONTRATO] AS varchar(30)) LIKE :filtro_contrato")
        params['filtro_contrato'] = f'%{filtro_contrato}%'
    if filtro_nome:
        condicoes.append("UPPER(c.[NO_MUTUARIO]) LIKE :filtro_nome")
        params['filtro_nome'] = f'%{filtro_nome.upper()}%'
    where = ('WHERE ' + ' AND '.join(condicoes)) if condicoes else ''

    rows = db.session.execute(text(f"""
        SELECT f.[DT_CALCULO],
               f.[NU_CONTRATO],
               MAX(c.[NO_MUTUARIO])                 AS NO_MUTUARIO,
               MAX(c.[MEMO])                        AS MEMO,
               MAX(c.[MEMO_GEFIN])                  AS MEMO_GEFIN,
               MAX(c.[GERENTE])                     AS GERENTE,
               MAX(c.[VR_LAUDO_AVALIACAO])          AS VR_LAUDO_AVALIACAO,
               -SUM(f.[VR_DEB_PROPTERREM])          AS VR_DEBITOS_PROPTERREM,
               MIN(f.[ANO_MES])                     AS DT_REFERENCIA,
               MAX(f.[NU_MES])                      AS ULTIMO_MES,
               SUM(f.[VR_PRESENTE])                 AS VPL,
               MAX(s.[IC_SITUACAO])                 AS IC_SITUACAO,
               MAX(s.[VR_VPL_SEM_PROPTERREM])       AS VR_VPL_SEM_PROPTERREM,
               MAX(s.[VR_DEB_PROPTERREM_SIMULADO])  AS VR_DEB_PROPTERREM_SIMULADO,
               MAX(s.[VR_REDUCAO_PROPTERREM])       AS VR_REDUCAO_PROPTERREM,
               MAX(s.[VR_VPL_SIMULADO])             AS VR_VPL_SIMULADO,
               MAX(s.[NU_MES_CONSOLIDACAO])         AS NU_MES_CONSOLIDACAO,
               MAX(s.[VR_TOTAL_SIMULADO])           AS VR_TOTAL_SIMULADO,
               MAX(s.[VR_PRESENTE_SIMULADO])        AS VR_PRESENTE_SIMULADO
        FROM {TB_FLUXOS} f
        LEFT JOIN {TB_CONTRATOS} c
               ON c.[NU_CONTRATO] = f.[NU_CONTRATO]
        LEFT JOIN {TB_SIMULACAO} s
               ON s.[DT_CALCULO] = f.[DT_CALCULO] AND s.[NU_CONTRATO] = f.[NU_CONTRATO]
        {where}
        GROUP BY f.[DT_CALCULO], f.[NU_CONTRATO]
        ORDER BY f.[DT_CALCULO] DESC, MAX(c.[NO_MUTUARIO]), f.[NU_CONTRATO]
    """), params).mappings().all()

    resumo = []
    for r in rows:
        vpl = _dec(r['VPL'])
        nu = contrato_texto(r['NU_CONTRATO'])
        situacao = r['IC_SITUACAO']
        if situacao is None:
            # Cálculo antigo sem simulação gravada: classifica só pelo sinal
            situacao = SITUACAO_POSITIVO if vpl >= 0 else None
        else:
            situacao = int(situacao)

        resumo.append({
            'dt_calculo': _para_date(r['DT_CALCULO']),
            'nu_contrato': nu,
            'nu_contrato_fmt': formatar_contrato(nu),
            'no_mutuario': r['NO_MUTUARIO'] or '(contrato não está mais cadastrado)',
            'memo': formatar_memo(r['MEMO']),
            # Dados da nota técnica (preenchidos no modal; NULL até lá)
            'memo_gefin': (r['MEMO_GEFIN'] or '').strip(),
            'gerente': (r['GERENTE'] or '').strip(),
            'vr_laudo': r['VR_LAUDO_AVALIACAO'],
            'vr_debitos': r['VR_DEBITOS_PROPTERREM'],
            'dt_referencia': _para_date(r['DT_REFERENCIA']),
            'ultimo_mes': int(r['ULTIMO_MES']),
            'vpl': vpl,
            'positivo': vpl >= 0,
            'situacao': situacao,
            'simulacao': {
                'VR_VPL_SEM_PROPTERREM': r['VR_VPL_SEM_PROPTERREM'],
                'VR_DEB_PROPTERREM_SIMULADO': r['VR_DEB_PROPTERREM_SIMULADO'],
                'VR_REDUCAO_PROPTERREM': r['VR_REDUCAO_PROPTERREM'],
                'VR_VPL_SIMULADO': r['VR_VPL_SIMULADO'],
                'NU_MES_CONSOLIDACAO': r['NU_MES_CONSOLIDACAO'],
                'VR_TOTAL_SIMULADO': r['VR_TOTAL_SIMULADO'],
                'VR_PRESENTE_SIMULADO': r['VR_PRESENTE_SIMULADO'],
            },
        })
    return resumo


def obter_fluxo_gravado(dt_calculo, nu_contrato, taxas=None):
    """Linhas gravadas do contrato + taxa da view (para exibição)."""
    rows = db.session.execute(text(f"""
        SELECT [NU_MES], [ANO_MES],
               [VR_DESP_MANUT], [VR_CUSTO_MANUT], [VR_DESP_CONSOL_PROP],
               [VR_DEB_PROPTERREM], [VR_VENDA], [VR_TOTAL], [VR_PRESENTE]
        FROM {TB_FLUXOS}
        WHERE [DT_CALCULO] = :dt_calculo
          AND [NU_CONTRATO] = :nu
        ORDER BY [NU_MES]
    """), {'dt_calculo': dt_calculo, 'nu': nu_contrato}).mappings().all()

    if taxas is None:
        taxas = obter_taxas_custo_medio()

    linhas = []
    taxa_anterior = None
    for r in rows:
        ano_mes = _para_date(r['ANO_MES'])
        taxa = taxas.get(chave_ano_mes(ano_mes)) if ano_mes else None
        if taxa is None:
            taxa = taxa_anterior
        taxa_anterior = taxa
        linha = dict(r)
        linha['ANO_MES'] = ano_mes
        linha['TAXA_AA'] = taxa
        linhas.append(linha)
    return linhas


# =========================================================================
# EXPORTAÇÃO EXCEL
# =========================================================================
_AZUL = '224ABE'
_CINZA = 'F8F9FC'
_FMT_VALOR = '#,##0.00;[Red]-#,##0.00'
_FMT_TAXA = '0.0000'


def _estilo_cabecalho(cel):
    cel.font = Font(bold=True, color='FFFFFF')
    cel.fill = PatternFill('solid', fgColor=_AZUL)
    cel.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cel.border = Border(bottom=Side(style='thin', color='B7C3E8'))


def _titulo(ws, linha, texto):
    cel = ws.cell(row=linha, column=1, value=texto)
    cel.font = Font(bold=True, size=13, color=_AZUL)


def _rotulo_valor(ws, linha, rotulo, valor, formato=None):
    c1 = ws.cell(row=linha, column=1, value=rotulo)
    c1.font = Font(bold=True)
    c2 = ws.cell(row=linha, column=2, value=valor)
    if formato:
        c2.number_format = formato
    return c2


def _float_ou_none(valor):
    return float(valor) if valor is not None else None


def _aba_fluxo(wb, nome_aba, titulo, cabecalho_info, linhas):
    """Cria uma aba com o fluxo mês a mês no layout da planilha."""
    cab_fluxo = ['Mês', 'Ano/Mês'] + [t for _, t in COLUNAS_FLUXO[:-1]] + \
                ['Custo Médio (% a.a.)', 'Valor Presente']
    larg_fluxo = [8, 12, 20, 18, 22, 20, 20, 18, 16, 20]

    wf = wb.create_sheet(nome_aba)
    _titulo(wf, 1, titulo)
    lin_info = 2
    for rot, val, fmt, cor in cabecalho_info:
        c = _rotulo_valor(wf, lin_info, rot, val, fmt)
        if cor:
            c.font = Font(bold=True, color=cor)
        lin_info += 1

    lin_cab = lin_info + 1
    for i, (t, w) in enumerate(zip(cab_fluxo, larg_fluxo), start=1):
        _estilo_cabecalho(wf.cell(row=lin_cab, column=i, value=t))
        wf.column_dimensions[get_column_letter(i)].width = w
    wf.row_dimensions[lin_cab].height = 45

    lin = lin_cab + 1
    primeira = lin
    for ln in linhas:
        wf.cell(row=lin, column=1, value=int(ln['NU_MES']))
        wf.cell(row=lin, column=2, value=ln['ANO_MES']).number_format = 'MMM/YYYY'
        col = 3
        for chave, _ in COLUNAS_FLUXO[:-1]:
            wf.cell(row=lin, column=col, value=float(ln[chave])).number_format = _FMT_VALOR
            col += 1
        wf.cell(row=lin, column=col, value=_float_ou_none(ln.get('TAXA_AA'))).number_format = _FMT_TAXA
        wf.cell(row=lin, column=col + 1, value=float(ln['VR_PRESENTE'])).number_format = _FMT_VALOR
        if lin % 2 == 0:
            for c in range(1, len(cab_fluxo) + 1):
                wf.cell(row=lin, column=c).fill = PatternFill('solid', fgColor=_CINZA)
        lin += 1

    tot = wf.cell(row=lin, column=1, value='TOTAL')
    tot.font = Font(bold=True)
    for c in list(range(3, 3 + len(COLUNAS_FLUXO) - 1)) + [len(cab_fluxo)]:
        letra = get_column_letter(c)
        cel = wf.cell(row=lin, column=c, value=f'=SUM({letra}{primeira}:{letra}{lin - 1})')
        cel.number_format = _FMT_VALOR
        cel.font = Font(bold=True)
    wf.freeze_panes = f'C{primeira}'
    return wf


def _nome_aba_unico(base, usados):
    nome = base[:31]
    k = 2
    while nome in usados:
        nome = f'{base[:28]}_{k}'
        k += 1
    usados.add(nome)
    return nome


def gerar_excel_calculo(resumo, fluxos):
    """
    Excel do fluxo: SOMENTE as abas de fluxo de cada contrato e, quando o
    contrato tem simulação (VPL negativo), a aba da Meta logo depois.
    (Sem abas de Resumo e de Parâmetros, a pedido da área.)

    resumo: lista de listar_resumo_calculo() (pode ter várias datas)
    fluxos: {(DT_CALCULO, nu_contrato str): linhas de obter_fluxo_gravado()}
    Retorna BytesIO com o .xlsx.
    """
    wb = Workbook()
    wb.remove(wb.active)   # tira a aba vazia padrão: a 1ª aba passa a ser o 1º fluxo

    # ---------------- Fluxo (e Meta, se negativo) por contrato ----------------
    usados = set()
    varias_datas = len({r['dt_calculo'] for r in resumo}) > 1
    for r in resumo:
        linhas = fluxos.get((r['dt_calculo'], r['nu_contrato']), [])
        cor_vpl = '047857' if r['positivo'] else 'B91C1C'
        info = SITUACOES.get(r['situacao'])
        sufixo_data = f" {r['dt_calculo'].strftime('%d%m%y')}" if varias_datas else ''

        _aba_fluxo(
            wb, _nome_aba_unico(f"Fluxo {r['nu_contrato']}{sufixo_data}", usados),
            f"Fluxo de Caixa Simulado - {r['no_mutuario']}",
            ([('Memo', r['memo'], None, None)] if r.get('memo') else []) + [
                ('Contrato', r['nu_contrato_fmt'], None, None),
                ('Data do cálculo', r['dt_calculo'], 'DD/MM/YYYY', None),
                ('Laudo de Avaliação', _float_ou_none(r['vr_laudo']), _FMT_VALOR, None),
                ('Débitos Propter Rem', _float_ou_none(r['vr_debitos']), _FMT_VALOR, None),
                ('Data de Referência (mês 0)', r['dt_referencia'], 'MM/YYYY', None),
                ('VP', float(r['vpl']), _FMT_VALOR, cor_vpl),
                ('Situação', info['rotulo'] if info else '-', None, info['cor'].lstrip('#') if info else None),
            ],
            linhas,
        )

        meta = montar_fluxo_meta(linhas, r['simulacao'])
        if meta:
            sim = r['simulacao']
            vpl_meta = sum((_dec(ln['VR_PRESENTE']) for ln in meta), Decimal('0'))
            _aba_fluxo(
                wb, _nome_aba_unico(f"Meta {r['nu_contrato']}{sufixo_data}", usados),
                f"Fluxo de Caixa Simulado - Meta - {r['no_mutuario']}",
                ([('Memo', r['memo'], None, None)] if r.get('memo') else []) + [
                    ('Contrato', r['nu_contrato_fmt'], None, None),
                    ('Situação', info['rotulo'] if info else '-', None,
                     info['cor'].lstrip('#') if info else None),
                    ('Débitos Propter Rem atuais', _float_ou_none(r['vr_debitos']), _FMT_VALOR, None),
                    ('Débitos Propter Rem simulados', _float_ou_none(sim['VR_DEB_PROPTERREM_SIMULADO']),
                     _FMT_VALOR, None),
                    ('Redução necessária', _float_ou_none(sim['VR_REDUCAO_PROPTERREM']), _FMT_VALOR, None),
                    ('VP da Meta', float(vpl_meta), _FMT_VALOR,
                     '047857' if vpl_meta >= 0 else 'B91C1C'),
                ],
                meta,
            )

    saida = BytesIO()
    wb.save(saida)
    saida.seek(0)
    return saida