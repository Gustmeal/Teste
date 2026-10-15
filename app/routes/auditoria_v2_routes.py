# -*- coding: utf-8 -*-
"""
Rotas: app/routes/auditoria_v2_routes.py
Telas da Auditoria V2 (somente administradores).

    /auditoria/v2/                    -> painel
    /auditoria/v2/atividades          -> tudo o que cada usuário fez (filtros + exportação Excel)
    /auditoria/v2/atividade/<id>      -> detalhe completo de uma ação (parâmetros, comandos, antes/depois)
    /auditoria/v2/alteracoes          -> todos os registros incluídos/alterados/excluídos
    /auditoria/v2/historico           -> linha do tempo de um registro específico
    /auditoria/v2/eventos             -> login, logout, login falho, erros, eventos de negócio
    /auditoria/v2/configuracoes       -> parâmetros e nomes dos aplicativos

Todas as opções de filtros e rótulos vêm do banco ou do backend.
Compatível com Python 3.9 e 3.12.
"""

import json
from datetime import datetime, timedelta
from io import BytesIO

from flask import Blueprint, flash, redirect, render_template, request, send_file, url_for
from flask_login import current_user, login_required
from sqlalchemy import and_, cast, func, or_, text

from app import db
from app.auth.utils import admin_required
from app.auditoria_v2 import config as aud_config
from app.auditoria_v2 import registrar_evento
from app.models.auditoria_v2 import (AudAlteracao, AudAlteracaoCampo, AudAplicativo, AudComandoSql,
                                     AudConfig, AudEvento, AudRequisicao)

auditoria_v2_bp = Blueprint('auditoria_v2', __name__, url_prefix='/auditoria/v2')

POR_PAGINA = 50

ROTULOS_TIPO = {
    'ACESSO': 'Acessou tela', 'CONSULTA': 'Consulta (dados da tela)', 'SUBMISSAO': 'Enviou formulário',
    'ALTERACAO': 'Alterou dados', 'EXPORTACAO': 'Exportou / baixou arquivo', 'LOGIN': 'Login',
    'LOGOUT': 'Logout', 'LOGIN_FALHO': 'Login com falha', 'ERRO': 'Erro no sistema',
    'ACESSO_NEGADO': 'Acesso negado',
}
CORES_TIPO = {
    'ACESSO': 'secondary', 'CONSULTA': 'info', 'SUBMISSAO': 'primary', 'ALTERACAO': 'warning',
    'EXPORTACAO': 'success', 'LOGIN': 'dark', 'LOGOUT': 'dark', 'LOGIN_FALHO': 'danger',
    'ERRO': 'danger', 'ACESSO_NEGADO': 'danger',
}
ROTULOS_OPERACAO = {'INSERT': 'Inclusão', 'UPDATE': 'Alteração', 'DELETE': 'Exclusão', 'MERGE': 'Mesclagem',
                    'PROCEDURE': 'Procedure', 'TRUNCATE': 'Limpeza total', 'DDL': 'Estrutura', 'LOTE': 'Lote SQL'}
CORES_OPERACAO = {'INSERT': 'success', 'UPDATE': 'warning', 'DELETE': 'danger', 'MERGE': 'info',
                  'PROCEDURE': 'secondary', 'TRUNCATE': 'danger', 'DDL': 'dark', 'LOTE': 'secondary'}
ROTULOS_SITUACAO = {'CONFIRMADA': 'Confirmada', 'DESFEITA': 'Desfeita (rollback)', 'ERRO': 'Erro'}
CORES_NIVEL = {'INFO': 'secondary', 'ALERTA': 'warning', 'CRITICO': 'danger'}
OPCOES_PERIODO = [(1, 'Hoje'), (7, 'Últimos 7 dias'), (30, 'Últimos 30 dias'), (90, 'Últimos 90 dias')]


@auditoria_v2_bp.context_processor
def _contexto_templates():
    return {
        'rotulos_tipo': ROTULOS_TIPO, 'cores_tipo': CORES_TIPO,
        'rotulos_operacao': ROTULOS_OPERACAO, 'cores_operacao': CORES_OPERACAO,
        'rotulos_situacao': ROTULOS_SITUACAO, 'cores_nivel': CORES_NIVEL,
        'current_year': datetime.now().year,
    }


# ======================================================================
# Auxiliares
# ======================================================================
def _eh_mssql():
    return db.engine.dialect.name == 'mssql'


def _dia(coluna):
    return cast(coluna, db.Date) if _eh_mssql() else func.date(coluna)


def _hora(coluna):
    return func.datepart(text('hour'), coluna) if _eh_mssql() else func.strftime('%H', coluna)


def _data_param(nome, padrao):
    valor = request.args.get(nome, '')
    for formato in ('%Y-%m-%d', '%d/%m/%Y', '%Y%m%d'):
        try:
            return datetime.strptime(valor, formato).date()
        except ValueError:
            continue
    return padrao


def _periodo_padrao(dias=7):
    hoje = datetime.now().date()
    ini = _data_param('dt_ini', hoje - timedelta(days=dias - 1))
    fim = _data_param('dt_fim', hoje)
    if ini > fim:
        ini, fim = fim, ini
    return ini, fim


def _intervalo(coluna, ini, fim):
    return and_(coluna >= datetime.combine(ini, datetime.min.time()),
                coluna < datetime.combine(fim + timedelta(days=1), datetime.min.time()))


def _json(texto):
    if not texto:
        return None
    try:
        return json.loads(texto)
    except Exception:
        return texto


def _usuarios_opcoes():
    linhas = db.session.query(AudRequisicao.USUARIO_ID, func.max(AudRequisicao.USUARIO_NOME)) \
        .filter(AudRequisicao.USUARIO_ID.isnot(None)) \
        .group_by(AudRequisicao.USUARIO_ID).all()
    return sorted([{'id': u, 'nome': n} for u, n in linhas], key=lambda x: (x['nome'] or '').lower())


def _aplicativos_opcoes():
    cadastrados = dict((a.BLUEPRINT, a.NO_APLICATIVO) for a in AudAplicativo.query.all())
    usados = [r[0] for r in db.session.query(AudRequisicao.BLUEPRINT).distinct().all() if r[0]]
    for bp in usados:
        cadastrados.setdefault(bp, aud_config.nome_padrao_blueprint(bp))
    return sorted([{'blueprint': b, 'nome': n} for b, n in cadastrados.items()], key=lambda x: x['nome'].lower())


def _query_atividades():
    ini, fim = _periodo_padrao(7)
    filtros = {
        'dt_ini': ini.strftime('%Y-%m-%d'), 'dt_fim': fim.strftime('%Y-%m-%d'),
        'usuario_id': request.args.get('usuario_id', ''), 'blueprint': request.args.get('blueprint', ''),
        'tipo': request.args.get('tipo', ''), 'busca': request.args.get('busca', '').strip(),
        'somente_alteracoes': request.args.get('somente_alteracoes', ''),
    }
    q = AudRequisicao.query.filter(_intervalo(AudRequisicao.DT_INICIO, ini, fim))
    if filtros['usuario_id']:
        try:
            q = q.filter(AudRequisicao.USUARIO_ID == int(filtros['usuario_id']))
        except ValueError:
            pass
    if filtros['blueprint']:
        q = q.filter(AudRequisicao.BLUEPRINT == filtros['blueprint'])
    if filtros['tipo']:
        q = q.filter(AudRequisicao.TIPO_ACAO == filtros['tipo'])
    if filtros['somente_alteracoes']:
        q = q.filter(AudRequisicao.QTDE_REGISTROS_ALTERADOS > 0)
    if filtros['busca']:
        termo = '%' + filtros['busca'] + '%'
        q = q.filter(or_(AudRequisicao.URL.like(termo), AudRequisicao.ROTA.like(termo),
                         AudRequisicao.PARAMETROS.like(termo), AudRequisicao.USUARIO_NOME.like(termo),
                         AudRequisicao.IP.like(termo), AudRequisicao.ARQUIVO_DOWNLOAD.like(termo)))
    return q, filtros


def _paginar(query, ordem):
    try:
        pagina = max(1, int(request.args.get('pagina', 1)))
    except ValueError:
        pagina = 1
    total = query.order_by(None).count()
    itens = query.order_by(*ordem).offset((pagina - 1) * POR_PAGINA).limit(POR_PAGINA).all()
    paginas = max(1, (total + POR_PAGINA - 1) // POR_PAGINA)
    return itens, {'pagina': pagina, 'paginas': paginas, 'total': total}


def _args_sem_pagina():
    return dict((k, v) for k, v in request.args.items() if k != 'pagina')


# ======================================================================
# Painel
# ======================================================================
@auditoria_v2_bp.route('/')
@login_required
@admin_required
def painel():
    try:
        dias = int(request.args.get('dias', 7))
    except ValueError:
        dias = 7
    if dias not in [d for d, _r in OPCOES_PERIODO]:
        dias = 7
    fim = datetime.now().date()
    ini = fim - timedelta(days=dias - 1)
    periodo_req = _intervalo(AudRequisicao.DT_INICIO, ini, fim)

    kpis = {
        'acoes': db.session.query(func.count(AudRequisicao.ID)).filter(periodo_req).scalar() or 0,
        'usuarios': db.session.query(func.count(func.distinct(AudRequisicao.USUARIO_ID)))
                              .filter(periodo_req).scalar() or 0,
        'alteracoes': db.session.query(func.count(AudAlteracao.ID))
                                .filter(_intervalo(AudAlteracao.DT_ALTERACAO, ini, fim)).scalar() or 0,
        'exportacoes': db.session.query(func.count(AudRequisicao.ID))
                                 .filter(periodo_req, AudRequisicao.TIPO_ACAO == 'EXPORTACAO').scalar() or 0,
        'erros': db.session.query(func.count(AudRequisicao.ID))
                           .filter(periodo_req, AudRequisicao.TIPO_ACAO == 'ERRO').scalar() or 0,
        'logins_falhos': db.session.query(func.count(AudEvento.ID))
                                   .filter(_intervalo(AudEvento.DT_EVENTO, ini, fim),
                                           AudEvento.TIPO == 'LOGIN_FALHO').scalar() or 0,
    }

    dia_col = _dia(AudRequisicao.DT_INICIO)
    por_dia = db.session.query(dia_col, func.count(AudRequisicao.ID)) \
        .filter(periodo_req).group_by(dia_col).order_by(dia_col).all()
    mapa_dia = dict((str(d)[:10], q) for d, q in por_dia)
    serie_dias = []
    d = ini
    while d <= fim:
        serie_dias.append({'dia': d.strftime('%d/%m'), 'qtde': mapa_dia.get(d.strftime('%Y-%m-%d'), 0)})
        d += timedelta(days=1)

    hora_col = _hora(AudRequisicao.DT_INICIO)
    por_hora_raw = db.session.query(hora_col, func.count(AudRequisicao.ID)) \
        .filter(periodo_req).group_by(hora_col).all()
    mapa_hora = dict((int(h), q) for h, q in por_hora_raw if h is not None)
    serie_horas = [{'hora': '{:02d}h'.format(h), 'qtde': mapa_hora.get(h, 0)} for h in range(24)]

    por_app = db.session.query(AudRequisicao.APLICATIVO, func.count(AudRequisicao.ID)) \
        .filter(periodo_req).group_by(AudRequisicao.APLICATIVO) \
        .order_by(func.count(AudRequisicao.ID).desc()).limit(12).all()
    por_tipo = db.session.query(AudRequisicao.TIPO_ACAO, func.count(AudRequisicao.ID)) \
        .filter(periodo_req).group_by(AudRequisicao.TIPO_ACAO) \
        .order_by(func.count(AudRequisicao.ID).desc()).all()
    por_usuario = db.session.query(
        AudRequisicao.USUARIO_ID, func.max(AudRequisicao.USUARIO_NOME), func.max(AudRequisicao.USUARIO_AREA),
        func.count(AudRequisicao.ID), func.sum(AudRequisicao.QTDE_REGISTROS_ALTERADOS),
        func.max(AudRequisicao.DT_INICIO)) \
        .filter(periodo_req, AudRequisicao.USUARIO_ID.isnot(None)) \
        .group_by(AudRequisicao.USUARIO_ID).order_by(func.count(AudRequisicao.ID).desc()).limit(15).all()
    tabelas = db.session.query(AudAlteracao.TABELA, AudAlteracao.OPERACAO, func.count(AudAlteracao.ID)) \
        .filter(_intervalo(AudAlteracao.DT_ALTERACAO, ini, fim)) \
        .group_by(AudAlteracao.TABELA, AudAlteracao.OPERACAO) \
        .order_by(func.count(AudAlteracao.ID).desc()).limit(15).all()
    alertas = AudEvento.query.filter(_intervalo(AudEvento.DT_EVENTO, ini, fim),
                                     AudEvento.NIVEL.in_(['ALERTA', 'CRITICO'])) \
        .order_by(AudEvento.DT_EVENTO.desc()).limit(15).all()

    graficos = {
        'dias': serie_dias,
        'horas': serie_horas,
        'apps': [{'nome': a or 'Não identificado', 'qtde': q} for a, q in por_app],
        'tipos': [{'nome': ROTULOS_TIPO.get(t, t), 'qtde': q} for t, q in por_tipo],
    }
    return render_template(
        'auditoria_v2/painel.html', dias=dias, opcoes_periodo=OPCOES_PERIODO, kpis=kpis,
        graficos=graficos, ini=ini, fim=fim,
        por_usuario=[{'id': u, 'nome': n, 'area': a, 'acoes': q, 'alterados': s or 0, 'ultima': ul}
                     for u, n, a, q, s, ul in por_usuario],
        tabelas=[{'tabela': t, 'operacao': o, 'qtde': q} for t, o, q in tabelas],
        alertas=alertas)


# ======================================================================
# Atividades
# ======================================================================
@auditoria_v2_bp.route('/atividades')
@login_required
@admin_required
def atividades():
    q, filtros = _query_atividades()
    itens, paginacao = _paginar(q, [AudRequisicao.DT_INICIO.desc(), AudRequisicao.ID.desc()])
    tipos = [r[0] for r in db.session.query(AudRequisicao.TIPO_ACAO).distinct().all() if r[0]]
    return render_template(
        'auditoria_v2/atividades.html', itens=itens, paginacao=paginacao, filtros=filtros,
        usuarios=_usuarios_opcoes(), aplicativos=_aplicativos_opcoes(),
        tipos=sorted(tipos, key=lambda t: ROTULOS_TIPO.get(t, t)), args=_args_sem_pagina())


@auditoria_v2_bp.route('/atividades/exportar')
@login_required
@admin_required
def exportar_atividades():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    q, _filtros = _query_atividades()
    linhas = q.order_by(AudRequisicao.DT_INICIO.desc()).limit(100000).all()
    wb = Workbook()
    ws = wb.active
    ws.title = 'Atividades'
    colunas = [('Data/Hora', 20), ('Usuário', 30), ('E-mail', 30), ('Área', 10), ('Aplicativo', 28),
               ('Ação', 24), ('Método', 8), ('Rota', 35), ('URL', 50), ('Status', 8), ('Registros alterados', 12),
               ('Arquivo baixado', 30), ('Duração (ms)', 12), ('IP', 16), ('Parâmetros', 60), ('Erro', 40),
               ('ID da ação', 34)]
    for i, (titulo, largura) in enumerate(colunas, start=1):
        cel = ws.cell(row=1, column=i, value=titulo)
        cel.font = Font(bold=True, color='FFFFFF')
        cel.fill = PatternFill('solid', fgColor='1F4E79')
        ws.column_dimensions[cel.column_letter].width = largura
    for r in linhas:
        ws.append([
            r.DT_INICIO.strftime('%d/%m/%Y %H:%M:%S') if r.DT_INICIO else '', r.USUARIO_NOME, r.USUARIO_EMAIL,
            r.USUARIO_AREA, r.APLICATIVO, ROTULOS_TIPO.get(r.TIPO_ACAO, r.TIPO_ACAO), r.METODO, r.ROTA, r.URL,
            r.STATUS_HTTP, r.QTDE_REGISTROS_ALTERADOS, r.ARQUIVO_DOWNLOAD, r.DURACAO_MS, r.IP,
            (r.PARAMETROS or '')[:32000], (r.ERRO or '')[:2000], r.ID_CORRELACAO,
        ])
    ws.freeze_panes = 'A2'
    saida = BytesIO()
    wb.save(saida)
    saida.seek(0)
    nome = 'auditoria_v2_atividades_{}.xlsx'.format(datetime.now().strftime('%Y%m%d_%H%M%S'))
    return send_file(saida, as_attachment=True, download_name=nome,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@auditoria_v2_bp.route('/atividade/<correlacao>')
@login_required
@admin_required
def atividade(correlacao):
    correlacao = (correlacao or '')[:32]
    requisicao = AudRequisicao.query.filter_by(ID_CORRELACAO=correlacao).first()
    comandos = AudComandoSql.query.filter_by(ID_CORRELACAO=correlacao).order_by(AudComandoSql.DT_EXECUCAO).all()
    alteracoes = AudAlteracao.query.filter_by(ID_CORRELACAO=correlacao).order_by(AudAlteracao.ID).all()
    eventos = AudEvento.query.filter_by(ID_CORRELACAO=correlacao).order_by(AudEvento.DT_EVENTO).all()
    if not (requisicao or comandos or alteracoes or eventos):
        flash('Ação de auditoria não encontrada.', 'warning')
        return redirect(url_for('auditoria_v2.atividades'))
    return render_template(
        'auditoria_v2/atividade.html', requisicao=requisicao, correlacao=correlacao,
        parametros=_json(requisicao.PARAMETROS) if requisicao else None,
        mensagens=_json(requisicao.MENSAGENS) if requisicao else None,
        comandos=comandos, eventos=eventos,
        alteracoes=[_montar_alteracao(a) for a in alteracoes])


def _montar_alteracao(a):
    antes = _json(a.DADOS_ANTES) or {}
    depois = _json(a.DADOS_DEPOIS) or {}
    if not isinstance(antes, dict):
        antes = {'valor': antes}
    if not isinstance(depois, dict):
        depois = {'valor': depois}
    campos = []
    for nome in list(dict.fromkeys(list(antes.keys()) + list(depois.keys()))):
        va = antes.get(nome)
        vd = depois.get(nome)
        campos.append({'campo': nome, 'antes': va, 'depois': vd,
                       'mudou': a.OPERACAO == 'UPDATE' and nome in antes and nome in depois and va != vd})
    if a.OPERACAO == 'UPDATE':
        campos.sort(key=lambda c: (not c['mudou'], c['campo']))
    return {'obj': a, 'campos': campos}


# ======================================================================
# Alterações de dados e histórico de registro
# ======================================================================
@auditoria_v2_bp.route('/alteracoes')
@login_required
@admin_required
def alteracoes():
    ini, fim = _periodo_padrao(7)
    filtros = {
        'dt_ini': ini.strftime('%Y-%m-%d'), 'dt_fim': fim.strftime('%Y-%m-%d'),
        'usuario_id': request.args.get('usuario_id', ''), 'blueprint': request.args.get('blueprint', ''),
        'tabela': request.args.get('tabela', ''), 'operacao': request.args.get('operacao', ''),
        'campo': request.args.get('campo', '').strip(), 'chave': request.args.get('chave', '').strip(),
        'valor': request.args.get('valor', '').strip(),
    }
    q = AudAlteracao.query.filter(_intervalo(AudAlteracao.DT_ALTERACAO, ini, fim))
    if filtros['usuario_id']:
        try:
            q = q.filter(AudAlteracao.USUARIO_ID == int(filtros['usuario_id']))
        except ValueError:
            pass
    if filtros['blueprint']:
        q = q.filter(AudAlteracao.BLUEPRINT == filtros['blueprint'])
    if filtros['tabela']:
        q = q.filter(AudAlteracao.TABELA == filtros['tabela'])
    if filtros['operacao']:
        q = q.filter(AudAlteracao.OPERACAO == filtros['operacao'])
    if filtros['chave']:
        q = q.filter(AudAlteracao.CHAVE_REGISTRO.like('%' + filtros['chave'] + '%'))
    if filtros['campo']:
        q = q.filter(AudAlteracao.CAMPOS_ALTERADOS.like('%' + filtros['campo'] + '%'))
    if filtros['valor']:
        termo = '%' + filtros['valor'] + '%'
        q = q.filter(or_(AudAlteracao.DADOS_ANTES.like(termo), AudAlteracao.DADOS_DEPOIS.like(termo)))
    itens, paginacao = _paginar(q, [AudAlteracao.DT_ALTERACAO.desc(), AudAlteracao.ID.desc()])
    tabelas = sorted(r[0] for r in db.session.query(AudAlteracao.TABELA).distinct().all() if r[0])
    operacoes = sorted(r[0] for r in db.session.query(AudAlteracao.OPERACAO).distinct().all() if r[0])
    return render_template(
        'auditoria_v2/alteracoes.html', itens=[_montar_alteracao(a) for a in itens], paginacao=paginacao,
        filtros=filtros, usuarios=_usuarios_opcoes(), aplicativos=_aplicativos_opcoes(), tabelas=tabelas,
        operacoes=operacoes, args=_args_sem_pagina())


@auditoria_v2_bp.route('/historico')
@login_required
@admin_required
def historico():
    tabela = request.args.get('tabela', '')
    chave = request.args.get('chave', '')
    if not tabela or not chave:
        flash('Informe a tabela e a chave do registro.', 'warning')
        return redirect(url_for('auditoria_v2.alteracoes'))
    versoes = AudAlteracao.query.filter_by(TABELA=tabela, CHAVE_REGISTRO=chave) \
        .order_by(AudAlteracao.DT_ALTERACAO.asc(), AudAlteracao.ID.asc()).all()
    return render_template('auditoria_v2/historico.html', tabela=tabela, chave=chave,
                           versoes=[_montar_alteracao(a) for a in versoes])


# ======================================================================
# Eventos
# ======================================================================
@auditoria_v2_bp.route('/eventos')
@login_required
@admin_required
def eventos():
    ini, fim = _periodo_padrao(7)
    filtros = {
        'dt_ini': ini.strftime('%Y-%m-%d'), 'dt_fim': fim.strftime('%Y-%m-%d'),
        'tipo': request.args.get('tipo', ''), 'nivel': request.args.get('nivel', ''),
        'usuario_id': request.args.get('usuario_id', ''), 'busca': request.args.get('busca', '').strip(),
    }
    q = AudEvento.query.filter(_intervalo(AudEvento.DT_EVENTO, ini, fim))
    if filtros['tipo']:
        q = q.filter(AudEvento.TIPO == filtros['tipo'])
    if filtros['nivel']:
        q = q.filter(AudEvento.NIVEL == filtros['nivel'])
    if filtros['usuario_id']:
        try:
            q = q.filter(AudEvento.USUARIO_ID == int(filtros['usuario_id']))
        except ValueError:
            pass
    if filtros['busca']:
        termo = '%' + filtros['busca'] + '%'
        q = q.filter(or_(AudEvento.DESCRICAO.like(termo), AudEvento.DETALHES.like(termo),
                         AudEvento.USUARIO_NOME.like(termo), AudEvento.ENTIDADE.like(termo)))
    itens, paginacao = _paginar(q, [AudEvento.DT_EVENTO.desc(), AudEvento.ID.desc()])
    tipos = sorted(r[0] for r in db.session.query(AudEvento.TIPO).distinct().all() if r[0])
    niveis = sorted(r[0] for r in db.session.query(AudEvento.NIVEL).distinct().all() if r[0])
    return render_template('auditoria_v2/eventos.html', itens=itens, paginacao=paginacao, filtros=filtros,
                           tipos=tipos, niveis=niveis, usuarios=_usuarios_opcoes(), args=_args_sem_pagina())


# ======================================================================
# Configurações
# ======================================================================
@auditoria_v2_bp.route('/configuracoes', methods=['GET', 'POST'])
@login_required
@admin_required
def configuracoes():
    if request.method == 'POST':
        try:
            agora = datetime.now()
            mudancas = []
            for cfg in AudConfig.query.all():
                campo = 'cfg_' + cfg.CHAVE
                if campo in request.form:
                    novo = request.form.get(campo, '').strip()
                    if novo != (cfg.VALOR or ''):
                        mudancas.append({'parametro': cfg.CHAVE, 'antes': cfg.VALOR, 'depois': novo})
                        cfg.VALOR = novo
                        cfg.UPDATED_AT = agora
                        cfg.UPDATED_POR = current_user.nome
            for app_aud in AudAplicativo.query.all():
                prefixo = 'app_{}_'.format(app_aud.ID)
                if (prefixo + 'nome') not in request.form:
                    continue
                nome = request.form.get(prefixo + 'nome', '').strip() or app_aud.NO_APLICATIVO
                categoria = request.form.get(prefixo + 'categoria', '').strip() or None
                auditar = request.form.get(prefixo + 'auditar') == '1'
                if (nome, categoria, auditar) != (app_aud.NO_APLICATIVO, app_aud.CATEGORIA, bool(app_aud.IN_AUDITAR)):
                    mudancas.append({'aplicativo': app_aud.BLUEPRINT,
                                     'antes': [app_aud.NO_APLICATIVO, app_aud.CATEGORIA, bool(app_aud.IN_AUDITAR)],
                                     'depois': [nome, categoria, auditar]})
                    app_aud.NO_APLICATIVO = nome
                    app_aud.CATEGORIA = categoria
                    app_aud.IN_AUDITAR = auditar
                    app_aud.UPDATED_AT = agora
            db.session.commit()
            aud_config.invalidar_cache()
            if mudancas:
                # A configuração da própria auditoria também fica auditada
                registrar_evento('CONFIG_AUDITORIA', 'Configuração da Auditoria V2 alterada',
                                 entidade='auditoria_v2', detalhes=mudancas, nivel='ALERTA')
            flash('Configurações da auditoria salvas. Valem em até 1 minuto para todos os usuários.', 'success')
        except Exception as erro:
            db.session.rollback()
            flash('Erro ao salvar configurações: {}'.format(erro), 'danger')
        return redirect(url_for('auditoria_v2.configuracoes'))

    desde = datetime.now() - timedelta(hours=24)
    rotas_frequentes = db.session.query(AudRequisicao.ROTA, func.count(AudRequisicao.ID)) \
        .filter(AudRequisicao.DT_INICIO >= desde, AudRequisicao.ROTA.isnot(None)) \
        .group_by(AudRequisicao.ROTA).order_by(func.count(AudRequisicao.ID).desc()).limit(15).all()
    volume = []
    for modelo, rotulo in ((AudRequisicao, 'Ações (requisições)'), (AudAlteracao, 'Registros alterados'),
                           (AudAlteracaoCampo, 'Campos alterados'), (AudComandoSql, 'Comandos SQL'),
                           (AudEvento, 'Eventos')):
        volume.append({'rotulo': rotulo, 'qtde': db.session.query(func.count(modelo.ID)).scalar() or 0})
    return render_template(
        'auditoria_v2/configuracoes.html', configs=AudConfig.query.order_by(AudConfig.CHAVE).all(),
        aplicativos=AudAplicativo.query.order_by(AudAplicativo.NO_APLICATIVO).all(),
        rotas_frequentes=[{'rota': r, 'qtde': q} for r, q in rotas_frequentes], volume=volume)
