# app/routes/analise_financeira_pf_routes.py
"""
Rotas do módulo Análise Financeira PF.

PRIMEIRA PARTE (esta versão):
  1. Parâmetros da análise  -> BDG.FIN_TB034_ANALISE_FINANCEIRA_PF_PARAMETROS
     - Exibe sempre a linha VIGENTE (DT_FIM_VIGENCIA IS NULL; se não houver,
       a de maior DT_INI_VIGENCIA).
     - Edição NÃO faz UPDATE nos valores: o usuário informa a
       DT_FIM_VIGENCIA da linha vigente, que é encerrada nessa data, e é
       INSERIDA uma nova linha com todos os valores repetidos, trocando
       apenas os editados, com DT_INI_VIGENCIA = DT_FIM informada + 1 dia
       (automático) e DT_FIM_VIGENCIA = NULL.
     - PZ_EXECUCAO_MESES é calculado no back end: PZ_EXECUCAO_ANOS * 12.

  2. Contratos da análise   -> BDG.FIN_TB035_ANALISE_FINANCEIRA_PF_CONTRATOS
     - Planilha-modelo gerada pelo próprio sistema (download).
     - Importação do Excel preenchido (arrastar e soltar): Contrato, Nome,
       Débitos Propter Rem, Laudo de Avaliação e Memo (opcional) — várias linhas.
     - Grava por NU_CONTRATO: se já existe atualiza, se não existe insere.
     - Permite excluir um contrato da lista.

SEGUNDA PARTE:
  3. Cálculo do fluxo       -> BDG.FIN_TB036_ANALISE_FINANCEIRA_PF_FLUXOS
     - Usuário escolhe os contratos (FIN_TB035) e a data de referência (mês 0).
     - Motor em app/utils/analise_financeira_pf_calculo.py reproduz a aba
       "Fluxo 1" da planilha (ver docstring do motor).
     - Taxa de desconto: CUSTO_MEDIO da view FIN_VW036 (% a.a.).
     - Parâmetros: versão da FIN_TB034 em vigor na data do cálculo.
     - Recalcular um contrato apaga o cálculo anterior dele (qualquer data).
  4. Resultados: consulta por data de cálculo, fluxo mês a mês (na própria
     página, abaixo da tabela) e exportação de toda a operação para Excel.
  5. VPL negativo -> Simulação de Equilíbrio (abas "Meta" da planilha),
     gravada em BDG.FIN_TB037_ANALISE_FINANCEIRA_PF_SIMULACAO:
       - Reversível: débitos propter rem simulados que zeram o VPL;
       - Irreversível: mesmo com débitos propter rem = 0 o VPL é negativo.
     Cada contrato tem um único cálculo: recalcular apaga o anterior.
  6. Nota Técnica em Word montada a partir da view FIN_VW037
     (app/utils/analise_financeira_pf_nota.py).
  7. Resumo por contrato a partir da view FIN_VW038: painel na tela de
     resultados (?painel=resumo) e exportação para Excel.

Todas as consultas usam text() parametrizado.
Compatível com Python 3.9 e 3.12.
"""
from flask import Blueprint, render_template, request, jsonify, send_file, url_for, flash, redirect
from flask_login import login_required, current_user
from app import db
from app.auth.decorators import sistema_requerido
from app.utils.audit import registrar_log
from datetime import datetime, date, timedelta
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from sqlalchemy import text
import re
import unicodedata
from io import BytesIO
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from app.utils import analise_financeira_pf_calculo as calc
from app.utils import analise_financeira_pf_nota as nota


analise_financeira_pf_bp = Blueprint(
    'analise_financeira_pf',
    __name__,
    url_prefix='/analise-financeira-pf'
)

SISTEMA = 'analise_financeira_pf'

TB_PARAMETROS = '[BDG].[FIN_TB034_ANALISE_FINANCEIRA_PF_PARAMETROS]'
TB_CONTRATOS = '[BDG].[FIN_TB035_ANALISE_FINANCEIRA_PF_CONTRATOS]'

# -------------------------------------------------------------------------
# Metadados das colunas de parâmetros (rótulo de tela + tipo de campo).
# Os VALORES sempre vêm do banco; aqui fica só a descrição de cada coluna.
#   tipo: 'valor'      -> R$, 2 casas
#         'percentual' -> exibido como está gravado na tabela, 2 casas
#         'prazo'      -> inteiro
#         'data'       -> date (dd/mm/aaaa na tela)
#         'texto'      -> varchar (limite em 'tamanho')
#         'numero'     -> inteiro (sem unidade)
#   obrigatorio: False -> pode ficar vazio (coluna aceita NULL); padrão True
#   editavel: False -> calculado no back end
# -------------------------------------------------------------------------
CAMPOS_PARAMETROS = [
    {'coluna': 'VR_CUSTO_CIPF', 'rotulo': 'Custo EMGEA Créditos Imobiliários PF (por contrato)',
     'tipo': 'valor', 'editavel': True},
    {'coluna': 'VR_CUSTO_INU', 'rotulo': 'Custo EMGEA Imóveis Não de Uso (por contrato)',
     'tipo': 'valor', 'editavel': True},
    {'coluna': 'VR_TARIFA_ADM_IMOVEIS', 'rotulo': 'Tarifa de Administração de Imóveis Não de Uso (por contrato)',
     'tipo': 'valor', 'editavel': True},
    {'coluna': 'PC_DESP_MANUT_INU_PRI_ANO', 'rotulo': 'Despesa de manutenção INU - primeiro ano',
     'tipo': 'percentual', 'editavel': True},
    {'coluna': 'PC_DESP_MANUT_INU_6_MESES', 'rotulo': 'Despesa de manutenção INU - seis meses subsequentes',
     'tipo': 'percentual', 'editavel': True},
    {'coluna': 'PC_DESP_DESCONTO_VENDA', 'rotulo': 'Desconto na venda do imóvel sobre o valor de avaliação',
     'tipo': 'percentual', 'editavel': True},
    {'coluna': 'VR_DESP_MEDIA_EXEC_JUD', 'rotulo': 'Despesa média de execução extrajudicial (por contrato)',
     'tipo': 'valor', 'editavel': True},
    {'coluna': 'PZ_EXECUCAO_ANOS', 'rotulo': 'Prazo para a execução (anos)',
     'tipo': 'prazo', 'editavel': True},
    {'coluna': 'PZ_PERMANENCIA_ESTOQUE_MESES', 'rotulo': 'Prazo de permanência em estoque (meses)',
     'tipo': 'prazo', 'editavel': True},
    {'coluna': 'PZ_EXECUCAO_MESES', 'rotulo': 'Prazo para a execução (meses) - calculado: anos × 12',
     'tipo': 'prazo', 'editavel': False},
    {'coluna': 'DT_CUSTO_DE_OPORTUNIDADE', 'rotulo': 'Data do custo de oportunidade',
     'tipo': 'data', 'editavel': True},
    # Dados de referência da Nota Técnica (opcionais: colunas aceitam NULL)
    {'coluna': 'NORMATIVO_SUFIN', 'rotulo': 'Normativo Sufin (ex.: FI.NOR.003.30)',
     'tipo': 'texto', 'editavel': True, 'obrigatorio': False, 'tamanho': 13},
    {'coluna': 'NR_ATA_DIREX', 'rotulo': 'Número da Ata Direx',
     'tipo': 'numero', 'editavel': True, 'obrigatorio': False},
    {'coluna': 'DT_ATA_DIREX', 'rotulo': 'Data da Ata Direx',
     'tipo': 'data', 'editavel': True, 'obrigatorio': False},
]

COLUNAS_VALORES = [c['coluna'] for c in CAMPOS_PARAMETROS]


@analise_financeira_pf_bp.context_processor
def inject_current_year():
    return {'current_year': datetime.utcnow().year}


# =========================================================================
# FUNÇÕES AUXILIARES
# =========================================================================
def _parse_decimal_br(valor):
    """
    Converte texto vindo do Excel/tela para Decimal.
    Aceita: 'R$ 2.026,04' | '2026,04' | '2026.04' | '223.000' | 223000 | ''
    Retorna None para vazio. Lança ValueError para texto inválido.
    """
    if valor is None:
        return None
    if isinstance(valor, (int, float, Decimal)):
        return Decimal(str(valor))

    txt = str(valor).strip()
    txt = txt.replace('R$', '').replace('%', '').replace('\u00a0', '').replace(' ', '')
    if txt == '' or txt == '-':
        return None

    negativo = False
    if txt.startswith('(') and txt.endswith(')'):
        negativo = True
        txt = txt[1:-1]

    if ',' in txt and '.' in txt:
        # padrão BR: 1.234.567,89
        txt = txt.replace('.', '').replace(',', '.')
    elif ',' in txt:
        # só vírgula: decimal BR
        txt = txt.replace(',', '.')
    elif '.' in txt:
        # só ponto: se for agrupamento de milhar (223.000 / 1.234.567) remove
        if re.fullmatch(r'-?\d{1,3}(\.\d{3})+', txt):
            txt = txt.replace('.', '')

    try:
        numero = Decimal(txt)
    except InvalidOperation:
        raise ValueError(f'Valor numérico inválido: "{valor}"')

    return -numero if negativo else numero


def _parse_inteiro(valor):
    """Converte para int (prazos). Retorna None para vazio."""
    numero = _parse_decimal_br(valor)
    if numero is None:
        return None
    if numero != numero.to_integral_value():
        raise ValueError(f'O prazo deve ser um número inteiro: "{valor}"')
    return int(numero)


def _parse_contrato(valor):
    """
    Limpa o número do contrato: aceita '1-0160-0101-095', '101600101095'.
    Recusa notação científica (ex.: '1,016E+11'), que é como o Excel copia
    a célula quando a coluna está com formato Geral — os últimos dígitos
    já se perderam nesse caso.
    """
    if valor is None:
        raise ValueError('Contrato não informado.')
    # Célula numérica do Excel (int / float / Decimal)
    if isinstance(valor, bool):
        raise ValueError(f'Contrato inválido: "{valor}".')
    if isinstance(valor, (int, float, Decimal)):
        numero = Decimal(str(valor))
        if numero != numero.to_integral_value() or numero <= 0:
            raise ValueError(f'Contrato inválido: "{valor}" (deve ser número inteiro).')
        if isinstance(valor, float) and numero >= Decimal('1e15'):
            raise ValueError(
                f'Contrato "{valor}" tem mais de 15 dígitos e o Excel perde precisão. '
                f'Formate a coluna Contrato como Texto.'
            )
        valor = str(int(numero))
    txt = str(valor).strip()
    if txt == '':
        raise ValueError('Contrato não informado.')
    if re.search(r'[eE][+\-]?\d+', txt):
        raise ValueError(
            f'Contrato "{txt}" está em notação científica. No Excel, formate a coluna '
            f'Contrato como Número sem casas decimais (ou Texto) e copie novamente.'
        )
    digitos = re.sub(r'[\s\.\-/]', '', txt)
    if not digitos.isdigit():
        raise ValueError(f'Contrato inválido: "{txt}" (use apenas números).')
    digitos = digitos.lstrip('0') or '0'
    # NU_CONTRATO é decimal(23,0): no máximo 23 dígitos
    if len(digitos) > 23:
        raise ValueError(f'Contrato "{txt}" tem mais de 23 dígitos.')
    return digitos


def _formatar_contrato(nu_contrato):
    """Contrato exibido só com os dígitos, sem traços: 101600101095."""
    return calc.formatar_contrato(nu_contrato)


def _fmt_br(valor, casas=2):
    """Formata número no padrão brasileiro."""
    if valor is None:
        return ''
    fmt = '{:,.' + str(casas) + 'f}'
    return fmt.format(float(valor)).replace(',', 'X').replace('.', ',').replace('X', '.')


def _formatar_parametro(valor, tipo):
    """Texto de exibição (tela, histórico)."""
    if valor is None:
        return ''
    if tipo == 'data':
        return _para_date(valor).strftime('%d/%m/%Y')
    if tipo == 'texto':
        return str(valor).strip()
    if tipo == 'valor':
        return _fmt_br(valor, 2)
    # VR_ e PC_ são decimal(18,2) no banco: 2 casas para ambos
    return _fmt_br(valor, 2) if tipo == 'percentual' else str(int(valor))


def _valor_input_parametro(valor, tipo):
    """Valor do campo de edição: data em ISO (input type=date), demais formatados."""
    if valor is None:
        return ''
    if tipo == 'data':
        return _para_date(valor).isoformat()
    return _formatar_parametro(valor, tipo)


def _parse_data_parametro(valor):
    """'2026-06-24' ou '24/06/2026' -> date(2026, 6, 24). Vazio -> None."""
    txt = str(valor or '').strip()
    if not txt:
        return None
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%Y%m%d'):
        try:
            return datetime.strptime(txt, fmt).date()
        except ValueError:
            continue
    raise ValueError(f'Data inválida: "{txt}".')


def _arredondar(valor, tipo):
    """Normaliza o valor para comparação/gravação conforme o tipo."""
    if valor is None:
        return None
    if tipo == 'data':
        return _para_date(valor)
    if tipo == 'texto':
        return str(valor).strip() or None
    if tipo in ('prazo', 'numero'):
        return int(valor)
    # VR_ e PC_ são decimal(18,2) no banco: arredonda em 2 casas
    return Decimal(str(valor)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)


def _obter_parametros_vigentes():
    """
    Retorna dict com a linha vigente de FIN_TB034 ou None.
    Regra: linha com DT_FIM_VIGENCIA IS NULL; se houver mais de uma (ou
    nenhuma), prevalece a de maior DT_INI_VIGENCIA.
    """
    sql = text(f"""
        SELECT TOP 1
               [DT_INI_VIGENCIA], [DT_FIM_VIGENCIA],
               [VR_CUSTO_CIPF], [VR_CUSTO_INU], [VR_TARIFA_ADM_IMOVEIS],
               [PC_DESP_MANUT_INU_PRI_ANO], [PC_DESP_MANUT_INU_6_MESES],
               [PC_DESP_DESCONTO_VENDA], [VR_DESP_MEDIA_EXEC_JUD],
               [PZ_EXECUCAO_ANOS], [PZ_PERMANENCIA_ESTOQUE_MESES],
               [PZ_EXECUCAO_MESES], [DT_CUSTO_DE_OPORTUNIDADE],
               [NORMATIVO_SUFIN], [NR_ATA_DIREX], [DT_ATA_DIREX]
        FROM {TB_PARAMETROS}
        ORDER BY CASE WHEN [DT_FIM_VIGENCIA] IS NULL THEN 0 ELSE 1 END,
                 [DT_INI_VIGENCIA] DESC
    """)
    row = db.session.execute(sql).mappings().first()
    return dict(row) if row else None


def _listar_historico_parametros():
    """Todas as vigências, da mais recente para a mais antiga."""
    sql = text(f"""
        SELECT [DT_INI_VIGENCIA], [DT_FIM_VIGENCIA],
               [VR_CUSTO_CIPF], [VR_CUSTO_INU], [VR_TARIFA_ADM_IMOVEIS],
               [PC_DESP_MANUT_INU_PRI_ANO], [PC_DESP_MANUT_INU_6_MESES],
               [PC_DESP_DESCONTO_VENDA], [VR_DESP_MEDIA_EXEC_JUD],
               [PZ_EXECUCAO_ANOS], [PZ_PERMANENCIA_ESTOQUE_MESES],
               [PZ_EXECUCAO_MESES], [DT_CUSTO_DE_OPORTUNIDADE],
               [NORMATIVO_SUFIN], [NR_ATA_DIREX], [DT_ATA_DIREX]
        FROM {TB_PARAMETROS}
        ORDER BY [DT_INI_VIGENCIA] DESC
    """)
    return [dict(r) for r in db.session.execute(sql).mappings().all()]


def _para_date(valor):
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.date()
    return valor


def _listar_contratos():
    sql = text(f"""
        SELECT [NU_CONTRATO], [NO_MUTUARIO],
               [VR_DEBITOS_PROPTERREM], [VR_LAUDO_AVALIACAO], [MEMO]
        FROM {TB_CONTRATOS}
        ORDER BY [NO_MUTUARIO], [NU_CONTRATO]
    """)
    contratos = []
    for r in db.session.execute(sql).mappings().all():
        nu = str(r['NU_CONTRATO']).strip()
        if nu.endswith('.0'):
            nu = nu[:-2]
        contratos.append({
            'nu_contrato': nu,
            'nu_contrato_fmt': _formatar_contrato(nu),
            'no_mutuario': r['NO_MUTUARIO'] or '',
            'vr_debitos': r['VR_DEBITOS_PROPTERREM'],
            'vr_debitos_fmt': _fmt_br(r['VR_DEBITOS_PROPTERREM'], 2),
            'vr_laudo': r['VR_LAUDO_AVALIACAO'],
            'vr_laudo_fmt': _fmt_br(r['VR_LAUDO_AVALIACAO'], 2),
            'memo': calc.formatar_memo(r['MEMO']),
        })
    return contratos


# =========================================================================
# PÁGINA PRINCIPAL
# =========================================================================
@analise_financeira_pf_bp.route('/')
@login_required
@sistema_requerido(SISTEMA)
def index():
    db.session.expire_all()

    vigente = _obter_parametros_vigentes()
    hoje = date.today()

    parametros = []
    for campo in CAMPOS_PARAMETROS:
        valor = vigente.get(campo['coluna']) if vigente else None
        parametros.append({
            'coluna': campo['coluna'],
            'rotulo': campo['rotulo'],
            'tipo': campo['tipo'],
            'editavel': campo['editavel'],
            # Tabela vazia: campos em branco para o primeiro cadastro
            'valor_fmt': _formatar_parametro(valor, campo['tipo']) if vigente else '',
            'valor_input': _valor_input_parametro(valor, campo['tipo']) if vigente else '',
        })

    dt_ini_vigente = _para_date(vigente['DT_INI_VIGENCIA']) if vigente else None
    # Linha mais recente ainda não começou: a anterior vale até o dia anterior
    vigencia_pendente = bool(dt_ini_vigente and dt_ini_vigente > hoje)
    dt_fim_anterior = (dt_ini_vigente - timedelta(days=1)) if vigencia_pendente else None

    historico = []
    for h in _listar_historico_parametros():
        historico.append({
            'dt_ini': _para_date(h['DT_INI_VIGENCIA']),
            'dt_fim': _para_date(h['DT_FIM_VIGENCIA']),
            'valores': [_formatar_parametro(h.get(c['coluna']), c['tipo']) for c in CAMPOS_PARAMETROS],
        })

    contratos = _listar_contratos()
    total_debitos = sum((c['vr_debitos'] or 0) for c in contratos)
    total_laudo = sum((c['vr_laudo'] or 0) for c in contratos)

    return render_template(
        'analise_financeira_pf/index.html',
        parametros=parametros,
        tem_parametros=vigente is not None,
        dt_ini_vigente=dt_ini_vigente,
        vigencia_pendente=vigencia_pendente,
        dt_fim_anterior=dt_fim_anterior,
        # Menor data de fim aceita: o próprio início da vigência atual
        dt_fim_minimo=dt_ini_vigente.isoformat() if dt_ini_vigente else '',
        campos_historico=CAMPOS_PARAMETROS,
        historico=historico,
        contratos=contratos,
        total_debitos_fmt=_fmt_br(total_debitos, 2),
        total_laudo_fmt=_fmt_br(total_laudo, 2),
    )


# =========================================================================
# PARÂMETROS — EDIÇÃO COM CONTROLE DE VIGÊNCIA
# =========================================================================
@analise_financeira_pf_bp.route('/parametros/salvar', methods=['POST'])
@login_required
@sistema_requerido(SISTEMA)
def salvar_parametros():
    """
    Recebe JSON:
      {"dt_fim_vigencia": "2026-12-31",
       "campos": {"VR_CUSTO_CIPF": "20,40", ...}}

    Lógica:
      - Compara cada campo recebido com a linha vigente (DT_FIM_VIGENCIA NULL).
      - Sem alteração de valor -> não grava nada.
      - Com alteração:
          a) DT_FIM_VIGENCIA é informada pelo usuário (obrigatória) e não
             pode ser anterior ao DT_INI_VIGENCIA da linha vigente.
          b) UPDATE na linha vigente: só DT_FIM_VIGENCIA = data informada.
          c) INSERT da nova linha: valores vigentes + editados,
             DT_INI_VIGENCIA = data informada + 1 dia (automático),
             DT_FIM_VIGENCIA = NULL.
      - Tabela vazia: INSERT da primeira linha com DT_INI_VIGENCIA = hoje.
    """
    try:
        dados = request.get_json(silent=True) or {}
        recebidos = dados.get('campos') or {}

        vigente = _obter_parametros_vigentes()
        hoje = date.today()

        # 1) Data de fim informada (só existe quando já há vigência)
        dt_fim = None
        if vigente:
            dt_fim = _parse_data_calculo(dados.get('dt_fim_vigencia'))
            if dt_fim is None:
                return jsonify({
                    'success': False,
                    'message': 'Informe a data de fim da vigência atual.'
                }), 400
            dt_ini_atual = _para_date(vigente['DT_INI_VIGENCIA'])
            if dt_fim < dt_ini_atual:
                return jsonify({
                    'success': False,
                    'message': ('A data de fim não pode ser anterior ao início da vigência atual ('
                                + dt_ini_atual.strftime('%d/%m/%Y') + ').')
                }), 400

        # 2) Monta os novos valores partindo da linha vigente
        novos = {}
        alterados = {}
        erros = []
        for campo in CAMPOS_PARAMETROS:
            col = campo['coluna']
            tipo = campo['tipo']
            atual = _arredondar(vigente.get(col), tipo) if vigente else None
            novos[col] = atual

            if not campo['editavel'] or col not in recebidos:
                continue
            try:
                if tipo in ('prazo', 'numero'):
                    valor = _parse_inteiro(recebidos.get(col))
                elif tipo == 'data':
                    valor = _parse_data_parametro(recebidos.get(col))
                elif tipo == 'texto':
                    valor = str(recebidos.get(col) or '').strip().upper() or None
                    if valor and len(valor) > campo.get('tamanho', 4000):
                        raise ValueError(f"máximo de {campo['tamanho']} caracteres.")
                else:
                    valor = _parse_decimal_br(recebidos.get(col))
            except ValueError as e:
                erros.append(f"{campo['rotulo']}: {e}")
                continue

            if valor is None:
                if campo.get('obrigatorio', True):
                    erros.append(f"{campo['rotulo']}: informe um valor.")
                    continue
                # Opcional deixado em branco: grava NULL
                if atual is not None:
                    alterados[col] = {'de': str(atual), 'para': None}
                novos[col] = None
                continue
            if tipo not in ('data', 'texto') and valor < 0:
                erros.append(f"{campo['rotulo']}: o valor não pode ser negativo.")
                continue

            valor = _arredondar(valor, tipo)
            if valor != atual:
                alterados[col] = {'de': None if atual is None else str(atual), 'para': str(valor)}
            novos[col] = valor

        if erros:
            return jsonify({'success': False, 'message': 'Corrija os campos: ' + ' | '.join(erros)}), 400

        # 3) Campo calculado: meses = anos * 12
        if novos.get('PZ_EXECUCAO_ANOS') is not None:
            meses = int(novos['PZ_EXECUCAO_ANOS']) * 12
            atual_meses = _arredondar(vigente.get('PZ_EXECUCAO_MESES'), 'prazo') if vigente else None
            if meses != atual_meses:
                alterados['PZ_EXECUCAO_MESES'] = {
                    'de': None if atual_meses is None else str(atual_meses), 'para': str(meses)
                }
            novos['PZ_EXECUCAO_MESES'] = meses

        if vigente and not alterados:
            return jsonify({'success': False, 'message': 'Nenhum valor foi alterado.'}), 400

        faltando = [c['rotulo'] for c in CAMPOS_PARAMETROS
                    if c.get('obrigatorio', True) and novos.get(c['coluna']) is None]
        if faltando:
            return jsonify({
                'success': False,
                'message': 'Preencha todos os parâmetros: ' + ', '.join(faltando)
            }), 400

        params_valores = {col: novos[col] for col in COLUNAS_VALORES}
        lista_colunas = ', '.join(f'[{c}]' for c in COLUNAS_VALORES)
        lista_params = ', '.join(f':{c}' for c in COLUNAS_VALORES)

        sql_insert = text(f"""
            INSERT INTO {TB_PARAMETROS}
                ([DT_INI_VIGENCIA], [DT_FIM_VIGENCIA], {lista_colunas})
            VALUES (:dt_ini, NULL, {lista_params})
        """)

        if not vigente:
            # Primeira linha da tabela
            db.session.execute(sql_insert, dict(params_valores, dt_ini=hoje))
            dt_ini_nova = hoje
            mensagem = 'Parâmetros cadastrados. Vigência a partir de ' + hoje.strftime('%d/%m/%Y') + '.'
        else:
            dt_ini_nova = dt_fim + timedelta(days=1)

            # Segurança: não pode já existir linha começando nessa data (PK)
            existe = db.session.execute(text(f"""
                SELECT COUNT(*) FROM {TB_PARAMETROS}
                WHERE [DT_INI_VIGENCIA] = :dt_ini_nova
            """), {'dt_ini_nova': dt_ini_nova}).scalar()
            if existe:
                return jsonify({
                    'success': False,
                    'message': ('Já existe uma vigência começando em '
                                + dt_ini_nova.strftime('%d/%m/%Y') + '. Escolha outra data de fim.')
                }), 400

            # Encerra a vigente com a data informada
            res = db.session.execute(text(f"""
                UPDATE {TB_PARAMETROS}
                SET [DT_FIM_VIGENCIA] = :dt_fim
                WHERE [DT_INI_VIGENCIA] = :dt_ini_atual
                  AND [DT_FIM_VIGENCIA] IS NULL
            """), {'dt_fim': dt_fim, 'dt_ini_atual': dt_ini_atual})
            if res.rowcount != 1:
                db.session.rollback()
                return jsonify({
                    'success': False,
                    'message': 'A vigência foi alterada por outra pessoa. Recarregue a página.'
                }), 409

            # Nova linha começa no dia seguinte ao fim informado
            db.session.execute(sql_insert, dict(params_valores, dt_ini=dt_ini_nova))
            mensagem = ('Vigência anterior encerrada em ' + dt_fim.strftime('%d/%m/%Y')
                        + '. Novos parâmetros valem a partir de ' + dt_ini_nova.strftime('%d/%m/%Y') + '.')

        db.session.commit()

        if vigente:
            log_antigos = {k: v['de'] for k, v in alterados.items()}
            log_antigos['DT_FIM_VIGENCIA'] = None
            log_novos = {k: v['para'] for k, v in alterados.items()}
            log_novos['DT_FIM_VIGENCIA'] = dt_fim.isoformat()
        else:
            log_antigos = None
            log_novos = {c: str(params_valores[c]) for c in COLUNAS_VALORES}

        registrar_log(
            acao='editar' if vigente else 'criar',
            entidade='analise_financeira_pf_parametros',
            entidade_id=0,
            descricao=f'Parâmetros Análise Financeira PF - vigência a partir de {dt_ini_nova.strftime("%d/%m/%Y")}',
            dados_antigos=log_antigos,
            dados_novos=log_novos,
        )

        return jsonify({'success': True, 'message': mensagem})

    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'Erro ao salvar parâmetros: {str(e)}'}), 500


# =========================================================================
# CONTRATOS — IMPORTAÇÃO POR EXCEL
# =========================================================================
# Cabeçalhos da planilha-modelo (a mesma lista é usada para gerar o modelo
# e para reconhecer as colunas na importação).
COLUNAS_MODELO = [
    {'chave': 'nu_contrato', 'titulo': 'Contrato', 'largura': 20, 'formato': '0'},
    {'chave': 'no_mutuario', 'titulo': 'Nome', 'largura': 45, 'formato': '@'},
    {'chave': 'vr_debitos', 'titulo': 'Débitos Propter Rem', 'largura': 22, 'formato': '#,##0.00'},
    {'chave': 'vr_laudo', 'titulo': 'Laudo de Avaliação', 'largura': 22, 'formato': '#,##0.00'},
    {'chave': 'memo', 'titulo': 'Memo SEI (número/ano)', 'largura': 26, 'formato': '@'},
]

# Colunas que precisam existir no cabeçalho (Memo é opcional: planilhas no
# modelo antigo, sem a coluna, continuam sendo aceitas)
COLUNAS_OBRIGATORIAS = {'nu_contrato', 'no_mutuario', 'vr_debitos', 'vr_laudo'}
# Memo: só "número/ano" (ex.: 239/2026). O aplicativo completa o restante.
REGEX_MEMO = re.compile(r'(\d{1,6})\s*/\s*(\d{4})')

EXTENSOES_PERMITIDAS = ('.xlsx', '.xlsm')
LINHAS_BUSCA_CABECALHO = 30   # procura o cabeçalho nas primeiras 30 linhas
MAX_LINHAS_VAZIAS = 20        # para de ler após 20 linhas vazias seguidas


def _normalizar_texto(valor):
    """'Débitos Propter Rem' -> 'debitos propter rem' (sem acento, minúsculo)."""
    if valor is None:
        return ''
    txt = unicodedata.normalize('NFKD', str(valor))
    txt = ''.join(ch for ch in txt if not unicodedata.combining(ch))
    return re.sub(r'\s+', ' ', txt).strip().lower()


def _identificar_coluna(cabecalho):
    """Diz qual campo um texto de cabeçalho representa (ou None)."""
    h = _normalizar_texto(cabecalho)
    if not h:
        return None
    # "Memo" vem antes: o texto do memorando não deve ser confundido com as outras colunas
    if h in ('memo', 'memorando', 'memo sei', 'memorando sei') or h.startswith('memo'):
        return 'memo'
    if 'contrato' in h:
        return 'nu_contrato'
    if 'nome' in h or 'mutuario' in h:
        return 'no_mutuario'
    if 'propter' in h or 'debito' in h:
        return 'vr_debitos'
    if 'laudo' in h or 'avaliacao' in h:
        return 'vr_laudo'
    return None


def _localizar_cabecalho(linhas):
    """
    Percorre as primeiras linhas e devolve (indice_linha, {campo: indice_coluna})
    da primeira linha que tenha as 4 colunas obrigatórias (Memo é opcional).
    Retorna (None, None) se não achar.
    """
    for idx, linha in enumerate(linhas):
        mapa = {}
        for col_idx, celula in enumerate(linha):
            campo = _identificar_coluna(celula)
            if campo and campo not in mapa:
                mapa[campo] = col_idx
        if COLUNAS_OBRIGATORIAS.issubset(mapa):
            return idx, mapa
    return None, None


def _validar_linha_contrato(bruto):
    """
    Valida e converte uma linha (dict com nu_contrato, no_mutuario,
    vr_debitos, vr_laudo). Retorna dict pronto para gravar ou lança ValueError.
    """
    nu = _parse_contrato(bruto.get('nu_contrato'))
    nome = str(bruto.get('no_mutuario') or '').strip()
    if not nome:
        raise ValueError('Nome do mutuário não informado.')
    if len(nome) > 100:
        raise ValueError('Nome do mutuário com mais de 100 caracteres.')
    debitos = _parse_decimal_br(bruto.get('vr_debitos'))
    laudo = _parse_decimal_br(bruto.get('vr_laudo'))
    if debitos is None:
        debitos = Decimal('0')
    if laudo is None:
        raise ValueError('Laudo de avaliação não informado.')
    if debitos < 0 or laudo < 0:
        raise ValueError('Valores não podem ser negativos.')

    # Memo: opcional; vazio vira None (o UPDATE mantém o memo já gravado)
    memo = _parse_memo(bruto.get('memo'))

    return {
        'nu': int(nu),   # NU_CONTRATO decimal(23,0) -> inteiro exato
        'nome': nome.upper(),
        'memo': memo,
        'debitos': debitos.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP),
        'laudo': laudo.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP),
    }


def _parse_memo(valor):
    """
    Extrai só o "número/ano" do memo: '239/2026' -> '239/2026'.
    Se colarem o texto completo ('Memorando SEI nº 239/2026/Gecoc/...'),
    também pega só o 239/2026. Vazio -> None. Fora do padrão -> erro.
    """
    txt = str(valor or '').strip()
    if not txt:
        return None
    m = REGEX_MEMO.search(txt)
    if not m:
        raise ValueError(f'Memo "{txt}" fora do padrão. Informe só número/ano, ex.: 239/2026.')
    numero, ano = int(m.group(1)), int(m.group(2))
    if numero <= 0 or not 1990 <= ano <= 2100:
        raise ValueError(f'Memo "{txt}" com número ou ano inválido.')
    return f'{numero}/{ano}'


def _ler_excel_contratos(conteudo):
    """
    Lê o Excel enviado e devolve (registros_validos: dict por contrato,
    erros: list[str], total_linhas: int).
    Prioriza a aba 'Contratos'; se não existir, usa a primeira aba.
    O número de linha nos erros é o número da linha no Excel.
    """
    wb = load_workbook(BytesIO(conteudo), read_only=True, data_only=True)
    try:
        nomes = {_normalizar_texto(n): n for n in wb.sheetnames}
        aba = wb[nomes['contratos']] if 'contratos' in nomes else wb[wb.sheetnames[0]]

        iterador = aba.iter_rows(values_only=True)

        # 1) Cabeçalho
        primeiras = []
        for _ in range(LINHAS_BUSCA_CABECALHO):
            try:
                primeiras.append(next(iterador))
            except StopIteration:
                break
        idx_cab, mapa = _localizar_cabecalho(primeiras)
        if idx_cab is None:
            raise ValueError(
                'Não encontrei o cabeçalho com as colunas Contrato, Nome, '
                'Débitos Propter Rem e Laudo de Avaliação. Use a planilha-modelo.'
            )

        # 2) Dados: o que sobrou das primeiras linhas + o restante da aba
        def _linhas_dados():
            for i, linha in enumerate(primeiras[idx_cab + 1:], start=idx_cab + 2):
                yield i, linha
            numero = len(primeiras) + 1
            for linha in iterador:
                yield numero, linha
                numero += 1

        validos = {}
        erros = []
        total = 0
        vazias_seguidas = 0
        for numero_linha, linha in _linhas_dados():
            linha = linha or ()
            bruto = {}
            for campo, col_idx in mapa.items():
                bruto[campo] = linha[col_idx] if col_idx < len(linha) else None

            if all(v is None or str(v).strip() == '' for v in bruto.values()):
                vazias_seguidas += 1
                if vazias_seguidas >= MAX_LINHAS_VAZIAS:
                    break
                continue
            vazias_seguidas = 0
            total += 1

            try:
                reg = _validar_linha_contrato(bruto)
                validos[reg['nu']] = reg   # contrato repetido: vale a última linha
            except ValueError as e:
                erros.append(f'Linha {numero_linha}: {e}')

        return validos, erros, total
    finally:
        wb.close()


def _gravar_contratos(validos):
    """UPDATE se o contrato existe, INSERT se não existe. Retorna (inseridos, atualizados)."""
    sql_update = text(f"""
        UPDATE {TB_CONTRATOS}
        SET [NO_MUTUARIO] = :nome,
            [VR_DEBITOS_PROPTERREM] = :debitos,
            [VR_LAUDO_AVALIACAO] = :laudo,
            [MEMO] = COALESCE(:memo, [MEMO])
        WHERE [NU_CONTRATO] = :nu
    """)
    sql_insert = text(f"""
        INSERT INTO {TB_CONTRATOS}
            ([NU_CONTRATO], [NO_MUTUARIO], [VR_DEBITOS_PROPTERREM], [VR_LAUDO_AVALIACAO], [MEMO])
        VALUES (:nu, :nome, :debitos, :laudo, :memo)
    """)
    inseridos = 0
    atualizados = 0
    for reg in validos.values():
        res = db.session.execute(sql_update, reg)
        if res.rowcount and res.rowcount > 0:
            atualizados += 1
        else:
            db.session.execute(sql_insert, reg)
            inseridos += 1
    return inseridos, atualizados


@analise_financeira_pf_bp.route('/contratos/modelo')
@login_required
@sistema_requerido(SISTEMA)
def baixar_modelo_contratos():
    """
    Gera a planilha-modelo para as áreas preencherem:
      - Aba 'Contratos': cabeçalho na linha 1, coluna Contrato com formato
        numérico sem casas (evita a notação científica 1,016E+11) e colunas
        de valor com formato monetário. Validação impede valor negativo.
      - Aba 'Instruções': como preencher.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = 'Contratos'

    fonte_cab = Font(bold=True, color='FFFFFF')
    fundo_cab = PatternFill('solid', fgColor='224ABE')
    borda = Border(bottom=Side(style='thin', color='B7C3E8'))
    centro = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for col_idx, col in enumerate(COLUNAS_MODELO, start=1):
        cel = ws.cell(row=1, column=col_idx, value=col['titulo'])
        cel.font = fonte_cab
        cel.fill = fundo_cab
        cel.alignment = centro
        cel.border = borda
        letra = get_column_letter(col_idx)
        ws.column_dimensions[letra].width = col['largura']
        # Formato das 1.000 primeiras linhas de dados
        for linha in range(2, 1002):
            ws.cell(row=linha, column=col_idx).number_format = col['formato']

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'A2'

    # Valores não negativos nas colunas de valor
    dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0',
                        allow_blank=True, showErrorMessage=True,
                        errorTitle='Valor inválido',
                        error='Informe um valor numérico maior ou igual a zero.')
    ws.add_data_validation(dv)
    dv.add('C2:D1001')

    # Contrato: número inteiro
    dv_ctr = DataValidation(type='whole', operator='greaterThan', formula1='0',
                            allow_blank=True, showErrorMessage=True,
                            errorTitle='Contrato inválido',
                            error='Informe o número do contrato só com dígitos (sem traços).')
    ws.add_data_validation(dv_ctr)
    dv_ctr.add('A2:A1001')

    # Aba de instruções
    wi = wb.create_sheet('Instruções')
    wi.column_dimensions['A'].width = 110
    instrucoes = [
        'Planilha-modelo - Análise Financeira PF (contratos)',
        '',
        '1. Preencha a aba "Contratos" a partir da linha 2, um contrato por linha.',
        '2. Contrato: somente números, sem traços (ex.: 101600101095).',
        '3. Nome: nome do mutuário (até 100 caracteres).',
        '4. Débitos Propter Rem: valor em reais. Deixe em branco ou 0 se não houver débitos.',
        '5. Laudo de Avaliação: valor do laudo em reais (obrigatório).',
        '6. Memo SEI (opcional): informe só o número e o ano do memorando, ex.: 239/2026.',
        '   O portal completa o restante: Memorando SEI nº 239/2026/Gecoc/Sucre/Diope.',
        '   Se ficar em branco, um contrato já cadastrado mantém o memo que já tinha.',
        '7. Não altere os títulos da linha 1 nem o nome da aba "Contratos".',
        '8. Contrato já cadastrado no portal é atualizado com os valores da planilha.',
        '9. Se o mesmo contrato aparecer mais de uma vez, vale a última linha.',
    ]
    for i, txt in enumerate(instrucoes, start=1):
        cel = wi.cell(row=i, column=1, value=txt)
        if i == 1:
            cel.font = Font(bold=True, size=13, color='224ABE')

    wb.active = 0

    saida = BytesIO()
    wb.save(saida)
    saida.seek(0)

    registrar_log(
        acao='exportar',
        entidade='analise_financeira_pf_contratos',
        entidade_id=0,
        descricao='Download da planilha-modelo de contratos da Análise Financeira PF',
    )

    return send_file(
        saida,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='Modelo_Contratos_Analise_Financeira_PF.xlsx',
    )


@analise_financeira_pf_bp.route('/contratos/importar', methods=['POST'])
@login_required
@sistema_requerido(SISTEMA)
def importar_contratos():
    """
    Recebe o Excel (campo 'arquivo', multipart/form-data).

    Lógica:
      1. Confere extensão (.xlsx / .xlsm).
      2. Procura o cabeçalho (Contrato, Nome, Débitos Propter Rem, Laudo de
         Avaliação) nas primeiras linhas da aba 'Contratos' (ou da 1ª aba).
         As colunas são achadas pelo título, então a ordem pode variar.
      3. Valida TODAS as linhas. Se alguma tiver erro, nada é gravado e a
         lista de erros volta com o número da linha do Excel.
      4. Sem erros: grava por NU_CONTRATO (atualiza se existe, insere se não).
    """
    try:
        arquivo = request.files.get('arquivo')
        if not arquivo or not arquivo.filename:
            return jsonify({'success': False, 'message': 'Selecione um arquivo Excel.'}), 400

        nome_arquivo = arquivo.filename
        if not nome_arquivo.lower().endswith(EXTENSOES_PERMITIDAS):
            return jsonify({
                'success': False,
                'message': 'Formato não aceito. Envie um arquivo .xlsx (use a planilha-modelo).'
            }), 400

        conteudo = arquivo.read()
        if not conteudo:
            return jsonify({'success': False, 'message': 'O arquivo está vazio.'}), 400

        try:
            validos, erros, total = _ler_excel_contratos(conteudo)
        except ValueError as e:
            return jsonify({'success': False, 'message': str(e)}), 400

        if erros:
            return jsonify({
                'success': False,
                'message': f'Nada foi gravado. {len(erros)} linha(s) com problema em {total} lida(s).',
                'erros': erros
            }), 400

        if not validos:
            return jsonify({'success': False, 'message': 'Nenhum contrato encontrado na planilha.'}), 400

        inseridos, atualizados = _gravar_contratos(validos)
        db.session.commit()

        registrar_log(
            acao='criar',
            entidade='analise_financeira_pf_contratos',
            entidade_id=0,
            descricao=(f'Importação de contratos Análise Financeira PF ({nome_arquivo}): '
                       f'{inseridos} novo(s), {atualizados} atualizado(s)'),
            dados_novos={'arquivo': nome_arquivo, 'contratos': [
                {'nu_contrato': str(r['nu']), 'no_mutuario': r['nome'],
                 'vr_debitos_propterrem': str(r['debitos']), 'vr_laudo_avaliacao': str(r['laudo']),
                 'memo': r['memo']}
                for r in validos.values()
            ]},
        )

        return jsonify({
            'success': True,
            'message': (f'Planilha importada: {inseridos} contrato(s) incluído(s) '
                        f'e {atualizados} atualizado(s).')
        })

    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'Erro ao importar planilha: {str(e)}'}), 500

@analise_financeira_pf_bp.route('/contratos/excluir', methods=['POST'])
@login_required
@sistema_requerido(SISTEMA)
def excluir_contrato():
    """Recebe JSON {"nu_contrato": "..."} e remove o contrato da lista."""
    try:
        dados = request.get_json(silent=True) or {}
        nu = int(_parse_contrato(dados.get('nu_contrato')))   # decimal(23,0)

        antigo = db.session.execute(text(f"""
            SELECT [NU_CONTRATO], [NO_MUTUARIO], [VR_DEBITOS_PROPTERREM], [VR_LAUDO_AVALIACAO], [MEMO]
            FROM {TB_CONTRATOS}
            WHERE [NU_CONTRATO] = :nu
        """), {'nu': nu}).mappings().first()

        if not antigo:
            return jsonify({'success': False, 'message': 'Contrato não encontrado.'}), 404

        db.session.execute(text(f"""
            DELETE FROM {TB_CONTRATOS}
            WHERE [NU_CONTRATO] = :nu
        """), {'nu': nu})
        db.session.commit()

        registrar_log(
            acao='excluir',
            entidade='analise_financeira_pf_contratos',
            entidade_id=0,
            descricao=f'Contrato {_formatar_contrato(nu)} excluído da Análise Financeira PF',
            dados_antigos={
                'nu_contrato': str(nu),
                'no_mutuario': antigo['NO_MUTUARIO'],
                'vr_debitos_propterrem': str(antigo['VR_DEBITOS_PROPTERREM']),
                'vr_laudo_avaliacao': str(antigo['VR_LAUDO_AVALIACAO']),
                'memo': antigo['MEMO'],
            },
        )

        return jsonify({'success': True, 'message': 'Contrato excluído.'})

    except ValueError as e:
        return jsonify({'success': False, 'message': str(e)}), 400
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'Erro ao excluir contrato: {str(e)}'}), 500


# =========================================================================
# CÁLCULO DO FLUXO — PÁGINA E EXECUÇÃO
# =========================================================================
CAMPOS_OBRIGATORIOS_CALCULO = [
    'VR_CUSTO_CIPF', 'VR_CUSTO_INU', 'VR_TARIFA_ADM_IMOVEIS',
    'PC_DESP_MANUT_INU_PRI_ANO', 'PC_DESP_MANUT_INU_6_MESES',
    'PC_DESP_DESCONTO_VENDA', 'VR_DESP_MEDIA_EXEC_JUD',
    'PZ_EXECUCAO_ANOS', 'PZ_PERMANENCIA_ESTOQUE_MESES',
]


def _parse_data_referencia(valor):
    """'2026-09' ou '2026-09-01' ou '09/2026' -> date(2026, 9, 1)."""
    txt = str(valor or '').strip()
    m = re.fullmatch(r'(\d{4})-(\d{2})(?:-\d{2})?', txt)
    if m:
        ano, mes = int(m.group(1)), int(m.group(2))
    else:
        m = re.fullmatch(r'(\d{2})/(\d{4})', txt)
        if not m:
            raise ValueError('Informe a data de referência (mês/ano).')
        mes, ano = int(m.group(1)), int(m.group(2))
    if not 1 <= mes <= 12:
        raise ValueError('Mês da data de referência inválido.')
    return date(ano, mes, 1)


def _parse_data_calculo(valor):
    try:
        return datetime.strptime(str(valor or '').strip(), '%Y-%m-%d').date()
    except ValueError:
        return None


@analise_financeira_pf_bp.route('/calculo')
@login_required
@sistema_requerido(SISTEMA)
def calculo():
    """Página de escolha dos contratos e da data de referência."""
    db.session.expire_all()
    hoje = date.today()

    parametros = calc.obter_parametros_na_data(hoje)

    contratos = []
    for c in calc.listar_contratos_cadastrados():
        nu = calc.contrato_texto(c['NU_CONTRATO'])
        contratos.append({
            'nu_contrato': nu,
            'nu_contrato_fmt': _formatar_contrato(nu),
            'no_mutuario': c['NO_MUTUARIO'] or '',
            'vr_debitos_fmt': _fmt_br(c['VR_DEBITOS_PROPTERREM'], 2),
            'vr_laudo_fmt': _fmt_br(c['VR_LAUDO_AVALIACAO'], 2),
        })

    # Faixa de meses que a view de custo médio cobre (limita o campo de data)
    taxas = calc.obter_taxas_custo_medio()
    mes_min = mes_max = None
    if taxas:
        k_min, k_max = min(taxas), max(taxas)
        mes_min = f'{k_min // 100:04d}-{k_min % 100:02d}'
        mes_max = f'{k_max // 100:04d}-{k_max % 100:02d}'

    return render_template(
        'analise_financeira_pf/calculo.html',
        contratos=contratos,
        tem_parametros=parametros is not None,
        mes_min=mes_min,
        mes_max=mes_max,
    )


@analise_financeira_pf_bp.route('/calculo/executar', methods=['POST'])
@login_required
@sistema_requerido(SISTEMA)
def executar_calculo():
    """
    Recebe JSON: {"contratos": ["101600101095", ...], "data_referencia": "2026-09"}

    Lógica:
      1. DT_CALCULO = hoje; parâmetros = versão da FIN_TB034 em vigor hoje.
      2. Taxas = FIN_VW036 (CUSTO_MEDIO por ANO_MES).
      3. Para cada contrato:
         a) calcula o fluxo (mês 0 = data de referência até
            PZ_EXECUCAO_MESES + PZ_PERMANENCIA_ESTOQUE_MESES);
         b) se o VPL for negativo, faz a simulação de equilíbrio (débitos
            propter rem que zeram o VPL, ou constata que nem zerando os
            débitos o VPL fica positivo);
         c) apaga TODO cálculo anterior do contrato (FIN_TB036 e FIN_TB037,
            qualquer data) e grava o novo — contrato repetido é refeito.
      4. Tudo em uma transação: se um contrato falhar, nada é gravado.
    """
    try:
        dados = request.get_json(silent=True) or {}

        try:
            data_ref = _parse_data_referencia(dados.get('data_referencia'))
        except ValueError as e:
            return jsonify({'success': False, 'message': str(e)}), 400

        lista = []
        for item in dados.get('contratos') or []:
            try:
                lista.append(int(_parse_contrato(item)))
            except ValueError as e:
                return jsonify({'success': False, 'message': str(e)}), 400
        lista = sorted(set(lista))
        if not lista:
            return jsonify({'success': False, 'message': 'Selecione ao menos um contrato.'}), 400

        dt_calculo = date.today()
        parametros = calc.obter_parametros_na_data(dt_calculo)
        if not parametros:
            return jsonify({
                'success': False,
                'message': 'Não há parâmetros em vigor hoje na FIN_TB034. Cadastre os parâmetros antes de calcular.'
            }), 400
        faltando = [c for c in CAMPOS_OBRIGATORIOS_CALCULO if parametros.get(c) is None]
        if faltando:
            return jsonify({
                'success': False,
                'message': 'Parâmetros incompletos na vigência atual: ' + ', '.join(faltando)
            }), 400

        taxas = calc.obter_taxas_custo_medio()
        if not taxas:
            return jsonify({'success': False, 'message': 'A view de custo médio (FIN_VW036) não retornou dados.'}), 400

        contratos = calc.obter_contratos(lista)
        encontrados = {int(_dec_para_int(c['NU_CONTRATO'])) for c in contratos}
        ausentes = [_formatar_contrato(n) for n in lista if n not in encontrados]
        if ausentes:
            return jsonify({
                'success': False,
                'message': 'Contrato(s) não cadastrado(s): ' + ', '.join(ausentes)
            }), 400

        mes_consolidacao, _ = calc.prazos_do_fluxo(parametros)

        resultados = []
        contagem = {s: 0 for s in calc.SITUACOES}
        for c in contratos:
            nu = int(_dec_para_int(c['NU_CONTRATO']))
            if c.get('VR_LAUDO_AVALIACAO') is None:
                raise calc.ErroCalculo(f'Contrato {_formatar_contrato(nu)} sem laudo de avaliação.')
            linhas, vpl = calc.calcular_fluxo(c, parametros, data_ref, taxas)
            simulacao = calc.simular_equilibrio(linhas, c.get('VR_DEBITOS_PROPTERREM'), mes_consolidacao)
            calc.gravar_fluxo(dt_calculo, nu, linhas, simulacao)
            contagem[simulacao['IC_SITUACAO']] += 1
            resultados.append({
                'nu_contrato': str(nu),
                'no_mutuario': c['NO_MUTUARIO'],
                'vpl': str(vpl),
                'situacao': simulacao['IC_SITUACAO'],
                'deb_propterrem_simulado': (None if simulacao['VR_DEB_PROPTERREM_SIMULADO'] is None
                                            else str(simulacao['VR_DEB_PROPTERREM_SIMULADO'])),
            })

        db.session.commit()

        registrar_log(
            acao='criar',
            entidade='analise_financeira_pf_fluxos',
            entidade_id=0,
            descricao=(f'Cálculo de fluxo Análise Financeira PF em {dt_calculo.strftime("%d/%m/%Y")}: '
                       f'{len(resultados)} contrato(s), referência {calc.mes_ano_texto(data_ref)}'),
            dados_novos={
                'dt_calculo': dt_calculo.isoformat(),
                'data_referencia': data_ref.isoformat(),
                'vigencia_parametros': str(_para_date(parametros['DT_INI_VIGENCIA'])),
                'contratos': resultados,
            },
        )

        return jsonify({
            'success': True,
            'message': (f'{len(resultados)} contrato(s) calculado(s): '
                        f'{contagem[calc.SITUACAO_POSITIVO]} positivo(s), '
                        f'{contagem[calc.SITUACAO_NEGATIVO_REVERSIVEL]} negativo(s) reversível(is), '
                        f'{contagem[calc.SITUACAO_NEGATIVO_IRREVERSIVEL]} negativo(s) irreversível(is).'),
            'redirect': url_for('analise_financeira_pf.resultados', dt_calculo=dt_calculo.isoformat()),
        })

    except calc.ErroCalculo as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': str(e)}), 400
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'Erro ao calcular: {str(e)}'}), 500


def _dec_para_int(valor):
    """Decimal('101600101095') / 101600101095.0 -> 101600101095"""
    return int(Decimal(str(valor)))


# =========================================================================
# RESULTADOS — CONSULTA, FLUXO E EXPORTAÇÃO
# =========================================================================
def _ler_filtros_resultados():
    """
    Lê os filtros da tela de resultados (query string):
      dt_calculo -> 'todos' (padrão) ou AAAA-MM-DD
      contrato   -> parte do número do contrato (traços/pontos são ignorados)
      nome       -> parte do nome do mutuário
    Devolve os valores prontos para a consulta e o dicionário 'url' com os
    filtros preenchidos, para manter os filtros nos links da página.
    """
    dt_txt = (request.args.get('dt_calculo') or 'todos').strip()
    dt_calculo = None if dt_txt == 'todos' else _parse_data_calculo(dt_txt)
    if dt_calculo is None:
        dt_txt = 'todos'

    contrato = re.sub(r'\D', '', request.args.get('contrato') or '')
    nome = re.sub(r'\s+', ' ', (request.args.get('nome') or '')).strip()

    url = {'dt_calculo': dt_txt}
    if contrato:
        url['contrato'] = contrato
    if nome:
        url['nome'] = nome

    return {
        'dt_calculo': dt_calculo,
        'dt_txt': dt_txt,
        'contrato': contrato,
        'nome': nome,
        'url': url,
        'ativos': bool(dt_calculo or contrato or nome),
    }


def _descricao_filtros(filtros):
    partes = []
    if filtros['dt_calculo']:
        partes.append('Data do cálculo ' + filtros['dt_calculo'].strftime('%d/%m/%Y'))
    else:
        partes.append('Todas as datas de cálculo')
    if filtros['contrato']:
        partes.append(f"Contrato contém {filtros['contrato']}")
    if filtros['nome']:
        partes.append(f"Nome contém {filtros['nome']}")
    return ' · '.join(partes)


@analise_financeira_pf_bp.route('/resultados')
@login_required
@sistema_requerido(SISTEMA)
def resultados():
    """
    Lista os cálculos gravados com VPL e situação de cada contrato.

    Filtros: data do cálculo (ou todas), parte do número do contrato e parte
    do nome. Com ?nu_contrato= e ?dt_fluxo= (ou quando o filtro retorna um
    único contrato) monta o fluxo mês a mês abaixo da tabela; ?visao=meta
    mostra o fluxo simulado (débitos propter rem que zeram o VPL).
    """
    db.session.expire_all()

    filtros = _ler_filtros_resultados()
    datas = calc.listar_datas_calculo()

    resumo = calc.listar_resumo_calculo(
        dt_calculo=filtros['dt_calculo'],
        filtro_contrato=filtros['contrato'] or None,
        filtro_nome=filtros['nome'] or None,
    )

    contagem = {s: 0 for s in calc.SITUACOES}
    for r in resumo:
        sim = r['simulacao']
        info = calc.SITUACOES.get(r['situacao'])
        r['dt_calculo_fmt'] = r['dt_calculo'].strftime('%d/%m/%Y') if r['dt_calculo'] else ''
        r['dt_calculo_iso'] = r['dt_calculo'].isoformat() if r['dt_calculo'] else ''
        r['vpl_fmt'] = _fmt_br(r['vpl'], 2)
        r['vr_laudo_fmt'] = _fmt_br(r['vr_laudo'], 2) if r['vr_laudo'] is not None else '-'
        r['vr_debitos_fmt'] = _fmt_br(r['vr_debitos'], 2)
        r['dt_referencia_fmt'] = calc.mes_ano_texto(r['dt_referencia'])
        r['situacao_info'] = info or {
            'rotulo': 'NEGATIVO', 'descricao': 'Recalcule o contrato para gerar a simulação.',
            'cor': '#b91c1c', 'fundo': '#fee2e2',
        }
        r['deb_simulado_fmt'] = (_fmt_br(sim['VR_DEB_PROPTERREM_SIMULADO'], 2)
                                 if sim['VR_DEB_PROPTERREM_SIMULADO'] is not None else '')
        r['reducao_fmt'] = (_fmt_br(sim['VR_REDUCAO_PROPTERREM'], 2)
                            if sim['VR_REDUCAO_PROPTERREM'] is not None else '')
        r['vpl_simulado_fmt'] = (_fmt_br(sim['VR_VPL_SIMULADO'], 2)
                                 if sim['VR_VPL_SIMULADO'] is not None else '')
        r['vpl_sem_fmt'] = (_fmt_br(sim['VR_VPL_SEM_PROPTERREM'], 2)
                            if sim['VR_VPL_SEM_PROPTERREM'] is not None else '')
        r['tem_meta'] = sim['VR_DEB_PROPTERREM_SIMULADO'] is not None
        r['memo'] = r.get('memo') or ''
        if r['situacao'] in contagem:
            contagem[r['situacao']] += 1

    total_vpl = sum((r['vpl'] for r in resumo), Decimal('0'))

    # Contrato selecionado para o detalhamento do fluxo (contrato + data do cálculo)
    contrato_sel = None
    nu_sel = None
    try:
        if request.args.get('nu_contrato'):
            nu_sel = _parse_contrato(request.args.get('nu_contrato'))
    except ValueError:
        nu_sel = None
    dt_fluxo = _parse_data_calculo(request.args.get('dt_fluxo'))

    if nu_sel is not None:
        contrato_sel = next((r for r in resumo if r['nu_contrato'] == nu_sel
                             and (dt_fluxo is None or r['dt_calculo'] == dt_fluxo)), None)
    elif len(resumo) == 1:
        contrato_sel = resumo[0]

    # Painel abaixo da tabela: 'fluxo' (padrão) ou 'resumo' (view FIN_VW038)
    painel = 'resumo' if request.args.get('painel') == 'resumo' else 'fluxo'

    visao = 'meta' if (request.args.get('visao') == 'meta' and contrato_sel
                       and contrato_sel['tem_meta']) else 'atual'
    fluxo = None
    resumo_contrato = None
    if contrato_sel and painel == 'fluxo':
        fluxo = _montar_fluxo_tela(contrato_sel['dt_calculo'], contrato_sel, visao)
    elif contrato_sel and painel == 'resumo':
        linhas_view = nota.carregar_resumo([int(contrato_sel['nu_contrato'])])
        resumo_contrato = nota.montar_resumo(linhas_view.get(contrato_sel['nu_contrato'], []),
                                             contrato_sel['no_mutuario'])

    return render_template(
        'analise_financeira_pf/resultados.html',
        datas=datas,
        filtros=filtros,
        resumo=resumo,
        total_vpl_fmt=_fmt_br(total_vpl, 2),
        total_vpl_positivo=total_vpl >= 0,
        qtd_positivos=contagem[calc.SITUACAO_POSITIVO],
        qtd_reversiveis=contagem[calc.SITUACAO_NEGATIVO_REVERSIVEL],
        qtd_irreversiveis=contagem[calc.SITUACAO_NEGATIVO_IRREVERSIVEL],
        contrato_sel=contrato_sel,
        fluxo=fluxo,
        visao=visao,
        painel=painel,
        resumo_contrato=resumo_contrato,
    )


def _montar_fluxo_tela(dt_calculo, contrato, visao='atual'):
    """
    Lê o fluxo gravado (ou monta a Meta, se visao='meta') e devolve as linhas
    já formatadas para a tela: cada célula com texto e indicador de negativo,
    rótulo do evento do mês (Consolidação / Manutenção / Venda) e totais.
    """
    linhas = calc.obter_fluxo_gravado(dt_calculo, int(contrato['nu_contrato']))
    if not linhas:
        return None
    if visao == 'meta':
        meta = calc.montar_fluxo_meta(linhas, contrato['simulacao'])
        if meta:
            linhas = meta
    mes_alterado = (int(contrato['simulacao']['NU_MES_CONSOLIDACAO'])
                    if visao == 'meta' and contrato['simulacao']['NU_MES_CONSOLIDACAO'] is not None else None)

    chaves = ['VR_DESP_MANUT', 'VR_CUSTO_MANUT', 'VR_DESP_CONSOL_PROP',
              'VR_DEB_PROPTERREM', 'VR_VENDA', 'VR_TOTAL']
    totais = {k: Decimal('0') for k in chaves + ['VR_PRESENTE']}
    saida = []

    for ln in linhas:
        valores = {k: Decimal(str(ln[k])) for k in chaves + ['VR_PRESENTE']}
        for k in totais:
            totais[k] += valores[k]

        eventos = []
        if valores['VR_DESP_CONSOL_PROP'] != 0 or int(ln['NU_MES']) == mes_alterado:
            eventos.append('Consolidação')
        if valores['VR_DESP_MANUT'] != 0:
            eventos.append('Manutenção')
        if valores['VR_VENDA'] != 0:
            eventos.append('Venda')

        saida.append({
            'nu_mes': int(ln['NU_MES']),
            'ano_mes': calc.mes_ano_texto(ln['ANO_MES']),
            'celulas': [{'txt': _fmt_br(valores[k], 2), 'neg': valores[k] < 0,
                         'simulado': (int(ln['NU_MES']) == mes_alterado
                                      and k in ('VR_DEB_PROPTERREM', 'VR_TOTAL'))}
                        for k in chaves],
            'taxa': _fmt_br(ln['TAXA_AA'], 4) if ln.get('TAXA_AA') is not None else '-',
            'vp': _fmt_br(valores['VR_PRESENTE'], 2),
            'vp_neg': valores['VR_PRESENTE'] < 0,
            'vp_simulado': int(ln['NU_MES']) == mes_alterado,
            'evento': ' + '.join(eventos),
        })

    return {
        'linhas': saida,
        'totais': [{'txt': _fmt_br(totais[k], 2), 'neg': totais[k] < 0} for k in chaves],
        'total_vp': _fmt_br(totais['VR_PRESENTE'], 2),
        'total_vp_neg': totais['VR_PRESENTE'] < 0,
    }


@analise_financeira_pf_bp.route('/resultados/exportar')
@login_required
@sistema_requerido(SISTEMA)
def exportar_resultados():
    """
    Exporta para Excel exatamente o que está filtrado na tela de resultados
    (data do cálculo ou todas, parte do contrato, parte do nome):
    somente a aba de fluxo de cada contrato e, quando houver, a aba da Meta.
    Com ?nu_contrato= e ?dt_fluxo= exporta só aquele contrato/cálculo.
    """
    filtros = _ler_filtros_resultados()

    nu_unico = None
    if request.args.get('nu_contrato'):
        try:
            nu_unico = int(_parse_contrato(request.args.get('nu_contrato')))
        except ValueError as e:
            return jsonify({'success': False, 'message': str(e)}), 400
    dt_fluxo = _parse_data_calculo(request.args.get('dt_fluxo'))

    if nu_unico is not None:
        resumo = calc.listar_resumo_calculo(dt_calculo=dt_fluxo or filtros['dt_calculo'],
                                            nu_contrato=nu_unico)
        descricao = f'Contrato {nu_unico}'
    else:
        resumo = calc.listar_resumo_calculo(
            dt_calculo=filtros['dt_calculo'],
            filtro_contrato=filtros['contrato'] or None,
            filtro_nome=filtros['nome'] or None,
        )
        descricao = _descricao_filtros(filtros)

    if not resumo:
        return jsonify({'success': False, 'message': 'Nenhum resultado para exportar.'}), 404

    taxas = calc.obter_taxas_custo_medio()
    fluxos = {(r['dt_calculo'], r['nu_contrato']):
              calc.obter_fluxo_gravado(r['dt_calculo'], int(r['nu_contrato']), taxas)
              for r in resumo}
    arquivo = calc.gerar_excel_calculo(resumo, fluxos)

    if nu_unico is not None:
        nome = f'Analise_Financeira_PF_{resumo[0]["nu_contrato"]}.xlsx'
    elif filtros['dt_calculo']:
        nome = f'Analise_Financeira_PF_{filtros["dt_calculo"].strftime("%Y%m%d")}.xlsx'
    else:
        nome = f'Analise_Financeira_PF_{date.today().strftime("%Y%m%d")}.xlsx'

    registrar_log(
        acao='exportar',
        entidade='analise_financeira_pf_fluxos',
        entidade_id=0,
        descricao=f'Exportação Excel Análise Financeira PF - {descricao} ({len(resumo)} contrato(s))',
    )

    return send_file(
        arquivo,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=nome,
    )


# =========================================================================
# NOTA TÉCNICA (WORD) — montada a partir da view FIN_VW037
# =========================================================================
@analise_financeira_pf_bp.route('/resultados/nota-tecnica')
@login_required
@sistema_requerido(SISTEMA)
def nota_tecnica():
    """
    Gera a Nota Técnica (Memorando) do contrato em Word.

    Lógica:
      1. Lê os fragmentos do contrato na view FIN_VW037 (ordem do ID).
      2. Troca cada '...' pela INFORMACAO formatada (R$, %, mês/ano, número);
         '...' sem valor fica destacado em amarelo para completar no Word.
      3. Junta os fragmentos em parágrafos e monta o .docx no layout dos
         modelos (Calibri 13,5, marcadores, itens numerados justificados).
      4. Acrescenta a assinatura do usuário logado (nome e cargo).
    Como o texto vem todo da view, a nota se ajusta sozinha ao resultado
    (positivo, negativo reversível ou irreversível).
    """
    filtros = _ler_filtros_resultados()
    try:
        nu = int(_parse_contrato(request.args.get('nu_contrato')))
    except ValueError as e:
        flash(str(e), 'danger')
        return redirect(url_for('analise_financeira_pf.resultados', **filtros['url']))

    registros = nota.carregar_textos(nu)
    if not registros:
        flash(f'A nota técnica do contrato {nu} ainda não está disponível na FIN_VW037.', 'warning')
        return redirect(url_for('analise_financeira_pf.resultados', **filtros['url']))

    paragrafos = nota.montar_paragrafos(registros)
    arquivo = nota.gerar_docx(
        paragrafos,
        assinante_nome=getattr(current_user, 'nome', None),
        assinante_cargo=getattr(current_user, 'cargo', None),
    )

    registrar_log(
        acao='exportar',
        entidade='analise_financeira_pf_nota',
        entidade_id=0,
        descricao=f'Nota técnica Análise Financeira PF gerada - contrato {nu}',
    )

    return send_file(
        arquivo,
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        as_attachment=True,
        download_name=f'Nota_Tecnica_AF_{nu}.docx',
    )


# =========================================================================
# RESUMO (FIN_VW038) — EXCEL
# =========================================================================
@analise_financeira_pf_bp.route('/resultados/resumo/exportar')
@login_required
@sistema_requerido(SISTEMA)
def exportar_resumo():
    """
    Exporta o resumo da view FIN_VW038 para Excel.
      - Com ?nu_contrato=: só aquele contrato.
      - Sem: todos os contratos do filtro atual da tela de resultados
        (aba 'Resumo' com todos + uma aba por contrato).
    """
    filtros = _ler_filtros_resultados()

    if request.args.get('nu_contrato'):
        try:
            nu = int(_parse_contrato(request.args.get('nu_contrato')))
        except ValueError as e:
            flash(str(e), 'danger')
            return redirect(url_for('analise_financeira_pf.resultados', **filtros['url']))
        resumo = calc.listar_resumo_calculo(nu_contrato=nu)
    else:
        resumo = calc.listar_resumo_calculo(
            dt_calculo=filtros['dt_calculo'],
            filtro_contrato=filtros['contrato'] or None,
            filtro_nome=filtros['nome'] or None,
        )

    # Um contrato aparece uma vez só (vale o cálculo mais recente)
    contratos = []
    vistos = set()
    for r in resumo:
        if r['nu_contrato'] not in vistos:
            vistos.add(r['nu_contrato'])
            contratos.append(r)

    if not contratos:
        flash('Nenhum contrato para exportar o resumo.', 'warning')
        return redirect(url_for('analise_financeira_pf.resultados', **filtros['url']))

    linhas_view = nota.carregar_resumo([int(r['nu_contrato']) for r in contratos])
    blocos = [{
        'nu_contrato': r['nu_contrato'],
        'no_mutuario': r['no_mutuario'],
        'resumo': nota.montar_resumo(linhas_view.get(r['nu_contrato'], []), r['no_mutuario']),
    } for r in contratos]

    arquivo = nota.gerar_excel_resumo(blocos)
    sufixo = contratos[0]['nu_contrato'] if len(contratos) == 1 else date.today().strftime('%Y%m%d')

    registrar_log(
        acao='exportar',
        entidade='analise_financeira_pf_resumo',
        entidade_id=0,
        descricao=f'Exportação do resumo Análise Financeira PF ({len(contratos)} contrato(s))',
    )

    return send_file(
        arquivo,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'Resumo_Analise_Financeira_PF_{sufixo}.xlsx',
    )